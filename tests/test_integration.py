"""End-to-end: RADIUS auth (PAP/CHAP/MS-CHAPv2) + accounting + NetFlow + web pages."""
import hashlib
import os
import socket
import struct
import tempfile
import time

_tmp = tempfile.mkdtemp()
os.environ.update({"DATA_DIR": _tmp, "RADIUS_BIND": "127.0.0.1", "RADIUS_AUTH_PORT": "18121",
                   "RADIUS_ACCT_PORT": "18131", "NETFLOW_BIND": "127.0.0.1", "NETFLOW_PORT": "20551",
                   "ADMIN_PASSWORD": "Secret12345"})

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pyrad import packet  # noqa: E402
from pyrad.client import Client  # noqa: E402
from pyrad.dictionary import Dictionary  # noqa: E402

from app import chap, config, radius_server  # noqa: E402
from app.db import (AcctSession, Group, Nas, SessionLocal, SpeedSample, TrafficCategory, User, encrypt,  # noqa: E402
                    init_db)
from app.netflow import Collector  # noqa: E402

SECRET = b"testing123"


@pytest.fixture(scope="module")
def env():
    init_db()
    db = SessionLocal()
    g = Group(name="کارمندان", rate_limit="1M/4M", quota_mb=100, quota_period="monthly", simultaneous_use=2)
    db.add(g)
    db.add(Nas(name="hq", ip="127.0.0.1", secret=SECRET.decode()))
    db.add(User(username="ali", password_enc=encrypt("Pass123"), group=g))
    db.add(User(username="night", password_enc=encrypt("x"), allowed_hours="fri 03:00-03:01"))
    db.commit()
    db.close()
    radius_server.start()
    time.sleep(0.3)
    cl = Client(server="127.0.0.1", authport=18121, acctport=18131, secret=SECRET,
                dict=Dictionary(config.RADIUS_DICTIONARY))
    cl.timeout = 3
    return cl


def test_pap_accept_with_attributes(env):
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="ali")
    req["User-Password"] = req.PwCrypt("Pass123")
    req["Calling-Station-Id"] = "3C:22:FB:11:22:33"
    req.add_message_authenticator()
    reply = env.SendPacket(req)
    assert reply.code == packet.AccessAccept
    assert reply["Mikrotik-Rate-Limit"][0] == "1M/4M"
    assert "Mikrotik-Total-Limit" not in reply      # quota counts internet only: the app enforces it, not the router
    assert reply.verify_message_authenticator(original_authenticator=req.authenticator)
    from app.db import set_setting
    db = SessionLocal()
    set_setting(db, "quota_internet_only", "0")
    db.commit()
    db.close()
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="ali")
    req["User-Password"] = req.PwCrypt("Pass123")
    req["Calling-Station-Id"] = "3C:22:FB:11:22:33"
    reply = env.SendPacket(req)
    assert reply["Mikrotik-Total-Limit"][0] == 100 * 1024 * 1024
    db = SessionLocal()
    set_setting(db, "quota_internet_only", "1")
    db.commit()
    db.close()


def test_pap_wrong_password(env):
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="ali")
    req["User-Password"] = req.PwCrypt("bad")
    assert env.SendPacket(req).code == packet.AccessReject


def test_hours_reject(env):
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="night")
    req["User-Password"] = req.PwCrypt("x")
    reply = env.SendPacket(req)
    assert reply.code == packet.AccessReject
    assert "ساعات" in reply["Reply-Message"][0]


def test_chap(env):
    challenge = os.urandom(16)
    ident = b"\x07"
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="ali")
    req["CHAP-Challenge"] = challenge
    req["CHAP-Password"] = ident + hashlib.md5(ident + b"Pass123" + challenge).digest()
    assert env.SendPacket(req).code == packet.AccessAccept


def test_mschapv2_vpn(env):
    ac, pc = os.urandom(16), os.urandom(16)
    nt = chap.generate_nt_response(ac, pc, "ali", "Pass123")
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="ali")
    req["Service-Type"] = 2
    req["MS-CHAP-Challenge"] = ac
    req["MS-CHAP2-Response"] = b"\x01\x00" + pc + b"\x00" * 8 + nt
    reply = env.SendPacket(req)
    assert reply.code == packet.AccessAccept
    expected = chap.authenticator_response("Pass123", nt, pc, ac, "ali")
    assert reply["MS-CHAP2-Success"][0] == b"\x01" + expected.encode()
    master = chap.mppe_master_key("Pass123", nt)
    assert _salt_decrypt(reply["MS-MPPE-Send-Key"][0], req.authenticator) == chap.mppe_asymmetric_key(master, True, True)


def _salt_decrypt(value, req_auth):
    """RFC 2548 decryption (pyrad's SaltDecrypt chains on plaintext, which is wrong for >1 block)."""
    salt, data, out, last = value[:2], value[2:], b"", req_auth + value[:2]
    while data:
        h = hashlib.md5(SECRET + last).digest()
        out += bytes(a ^ b for a, b in zip(h, data[:16]))
        last, data = data[:16], data[16:]
    return out[1:1 + out[0]]


def _acct(env, status, sid, i, o, t):
    req = env.CreateAcctPacket(User_Name="ali")
    req["Acct-Status-Type"] = status
    req["Acct-Session-Id"] = sid
    req["Framed-IP-Address"] = "10.5.50.10"
    req["Calling-Station-Id"] = "3C:22:FB:11:22:33"
    req["Acct-Input-Octets"] = i
    req["Acct-Output-Octets"] = o
    req["Acct-Session-Time"] = t
    return env.SendPacket(req)


def test_accounting_and_netflow_and_web(env):
    assert _acct(env, 1, "S1", 0, 0, 0).code == packet.AccountingResponse
    time.sleep(1.1)
    _acct(env, 3, "S1", 1_000_000, 50_000_000, 60)
    db = SessionLocal()
    s = db.query(AcctSession).filter_by(acct_session_id="S1").one()
    assert s.active and s.bytes_out == 50_000_000 and s.device_type == "mobile"
    assert db.query(SpeedSample).count() == 1
    db.close()

    # NetFlow v5: 10MB download from googlevideo-like Google IP + Telegram upload
    col = Collector()
    hdr = struct.pack("!HHIIIIBBH", 5, 2, 0, int(time.time()), 0, 0, 0, 0, 0)

    def rec(src, dst, octets, sport, dport, proto):
        return (socket.inet_aton(src) + socket.inet_aton(dst) + b"\0" * 8 + struct.pack("!II", 10, octets) +
                b"\0" * 8 + struct.pack("!HH", sport, dport) + b"\0" + b"\0" + bytes([proto]) + b"\0" * 9)
    pkt = hdr + rec("142.250.10.10", "10.5.50.10", 10 * 1024 * 1024, 443, 50000, 17) + \
        rec("10.5.50.10", "149.154.167.50", 2048, 50001, 443, 6)
    col.handle(pkt, "127.0.0.1")
    col.flush()
    db = SessionLocal()
    cats = {c.category: c for c in db.query(TrafficCategory).all()}
    assert cats["video"].bytes_down == 10 * 1024 * 1024
    assert cats["messaging"].bytes_up == 2048
    db.close()

    _acct(env, 2, "S1", 2_000_000, 60_000_000, 120)
    from app.web import app
    with TestClient(app) as c:
        assert c.get("/", follow_redirects=False).status_code == 303
        c.post("/login", data={"username": "admin", "password": "Secret12345"})
        for url in ["/", "/users", "/users/new", "/users/1/edit", "/groups", "/groups/new", "/nas", "/ad", "/live",
                    "/reports", "/reports?preset=month&user_id=1", "/devices", "/logs",
                    "/reports/export?kind=users", "/reports/export?kind=sessions", "/api/live"]:
            r = c.get(url)
            assert r.status_code == 200, url
        r = c.get("/reports?preset=today")
        assert "استریم ویدیو" in r.text and "ali" in r.text
        r = c.post("/users/save", data={"username": "reza", "password": "p", "enabled": "1",
                                        "allowed_hours": "sat-wed 08:00-17:00", "quota_value": "5",
                                        "quota_unit": "GB", "rate_limit": "2M/8M"}, follow_redirects=False)
        assert r.status_code == 303
        db = SessionLocal()
        u = db.query(User).filter_by(username="reza").one()
        assert u.quota_mb == 5120 and u.rate_limit == "2M/8M"
        db.close()
        assert c.get("/ua?mac=AA:BB:CC:00:11:22&ua=Mozilla/5.0 (Windows NT 10.0)").status_code == 200


def _login(c, user, pw, new_pw=None):
    r = c.post("/login", data={"username": user, "password": pw}, follow_redirects=False)
    if new_pw:
        assert r.headers["location"] == "/account"
        assert c.get("/users", follow_redirects=False).headers["location"] == "/account"
        c.post("/account", data={"current": pw, "new": new_pw, "again": new_pw})
    return r


def test_roles_and_management_reports(env):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.db import PanelUser
    from app.web import app
    with TestClient(app) as c:  # default IT / HR accounts are created disabled (their password is well known)
        r = c.post("/login", data={"username": "it", "password": "it12345678"}, follow_redirects=True)
        assert "غیرفعال" in r.text
    db = SessionLocal()
    for u in db.query(PanelUser).filter(PanelUser.username.in_(["it", "hr"])):
        u.enabled = True
    db.commit()
    db.close()
    with TestClient(app) as c:  # IT expert
        _login(c, "it", "it12345678", "ItPass2026x")
        assert c.get("/users").status_code == 200
        assert c.get("/live").status_code == 200
        for url in ["/nas", "/admins", "/mgmt", "/attendance", "/audit"]:
            assert c.get(url).status_code == 403, url
        assert c.post("/ad/save", data={}).status_code == 403
        assert "کاربران پنل" not in c.get("/users").text
    with TestClient(app) as c:  # HR expert: read-only reports
        _login(c, "hr", "hr12345678", "HrPass2026x")
        assert c.get("/mgmt").status_code == 200
        assert c.get("/mgmt/print?preset=month").status_code == 200
        r = c.get("/mgmt/export.xlsx?preset=month")
        assert r.status_code == 200
        wb = load_workbook(BytesIO(r.content))
        assert "خلاصه" in wb.sheetnames and "کاربران" in wb.sheetnames and wb["خلاصه"].sheet_view.rightToLeft
        assert c.get("/attendance?preset=30d").status_code == 200
        r = c.get("/attendance/export.xlsx?preset=30d")
        assert load_workbook(BytesIO(r.content)).sheetnames == ["ساعات اتصال", "خلاصه کارکنان"]
        assert c.get("/users").status_code == 200
        assert "ویرایش" not in c.get("/users").text
        for url in ["/users/new", "/live", "/nas", "/devices", "/admins"]:
            assert c.get(url).status_code == 403, url
        assert c.post("/users/save", data={"username": "x"}).status_code == 403
    with TestClient(app) as c:  # admin manages panel users and sees the audit log
        _login(c, "admin", "Secret12345")
        r = c.post("/admins/save", data={"username": "it2", "full_name": "کارشناس دوم", "role": "it",
                                         "password": "Abcdef123", "enabled": "1", "custom_perms": "1",
                                         "perm_users_view": "1", "perm_reports_view": "1"}, follow_redirects=False)
        assert r.status_code == 303
        assert "it2" in c.get("/admins").text
        audit = c.get("/audit").text
        assert "خروجی Excel گزارش مدیریتی" in audit and "ذخیره کاربر پنل" in audit
    with TestClient(app) as c:  # custom permissions are enforced
        _login(c, "it2", "Abcdef123")
        assert c.get("/reports").status_code == 200
        assert c.get("/live").status_code == 403
    with TestClient(app) as c:  # lockout after 5 failures
        for _ in range(5):
            _login(c, "it2", "wrong")
        r = c.post("/login", data={"username": "it2", "password": "Abcdef123"})
        assert "قفل" in r.text


def test_attendance_merging(env):
    from datetime import datetime

    from app import mgmt
    db = SessionLocal()
    u = db.query(User).filter_by(username="ali").one()
    day = datetime(2026, 10, 10)  # Saturday; default work hours 08:00-16:00
    db.add(AcctSession(acct_session_id="A1", nas_ip="127.0.0.1", username="ali", user_id=u.id, active=False,
                       start_time=day.replace(hour=8), stop_time=day.replace(hour=12), last_update=day.replace(hour=12)))
    db.add(AcctSession(acct_session_id="A2", nas_ip="127.0.0.1", username="ali", user_id=u.id, active=False,
                       start_time=day.replace(hour=11), stop_time=day.replace(hour=18), last_update=day.replace(hour=18)))
    db.commit()
    rows = mgmt.attendance(db, day, day.replace(hour=23, minute=59), ids=[u.id])
    assert len(rows) == 1
    r = rows[0]
    assert r["first"].hour == 8 and r["last"].hour == 18 and r["sessions"] == 2
    assert r["online"] == 10 * 3600 and r["inside"] == 8 * 3600 and r["outside"] == 2 * 3600
    db.close()


def test_offices_departments_and_unit_manager_scope(env):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.db import Department, Office, PanelUser
    from app.web import app
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        assert c.get("/org").status_code == 200
        c.post("/org/office/save", data={"name": "دفتر یزد", "city": "یزد"})
        c.post("/org/office/save", data={"name": "دفتر تهران", "city": "تهران"})
        db = SessionLocal()
        yazd = db.query(Office).filter_by(name="دفتر یزد").one()
        tehran = db.query(Office).filter_by(name="دفتر تهران").one()
        db.close()
        c.post("/org/dept/save", data={"name": "مالی", "office_id": ""})
        db = SessionLocal()
        fin = db.query(Department).filter_by(name="مالی").one()
        db.close()
        r = c.post("/users/save", data={"username": "yazd1", "password": "p1", "enabled": "1",
                                        "office_id": str(yazd.id), "department_id": str(fin.id),
                                        "personnel_code": "1001"}, follow_redirects=False)
        assert r.status_code == 303
        c.post("/users/save", data={"username": "teh1", "password": "p2", "enabled": "1", "office_id": str(tehran.id),
                                    "department_id": str(fin.id)})
        c.post("/org/assign", data={"usernames": "ali", "office_id": str(yazd.id), "department_id": "keep"})
        # unit manager without scope is refused
        r = c.post("/admins/save", data={"username": "mgr_yazd", "role": "unit_manager", "password": "Yazd12345",
                                         "enabled": "1"}, follow_redirects=False)
        assert r.headers["location"] == "/admins/new"
        c.post("/admins/save", data={"username": "mgr_yazd", "full_name": "مدیر دفتر یزد", "role": "unit_manager",
                                     "password": "Yazd12345", "enabled": "1", "scope_offices": str(yazd.id)})
        c.post("/admins/save", data={"username": "mgr_fin", "role": "unit_manager", "password": "Fin123456",
                                     "enabled": "1", "scope_departments": str(fin.id)})
        assert "مدیر دفتر یزد" in c.get("/org").text
        db = SessionLocal()
        assert db.query(PanelUser).filter_by(username="mgr_yazd").one().scope_offices == str(yazd.id)
        db.close()
    with TestClient(app) as c:  # office manager: only Yazd staff, view only
        _login(c, "mgr_yazd", "Yazd12345")
        t = c.get("/users").text
        assert "yazd1" in t and "ali" in t and "teh1" not in t and "reza" not in t
        assert "محدوده دسترسی شما" in t and "ویرایش" not in t
        assert "teh1" not in c.get("/reports?preset=month").text
        assert c.get("/mgmt?preset=month").status_code == 200
        wb = load_workbook(BytesIO(c.get("/mgmt/export.xlsx?preset=month").content))
        users = [row[1] for row in wb["کاربران"].iter_rows(min_row=2, values_only=True)]
        assert "teh1" not in users
        r = c.get(f"/reports?preset=month&office_id={tehran.id}")
        assert r.status_code == 200 and "teh1" not in r.text
        assert c.get("/attendance?preset=30d").status_code == 200
        assert c.get("/live").status_code == 200
        for url in ["/users/new", "/org", "/nas", "/admins", "/devices", "/logs", "/mt-local", "/hotspot"]:
            assert c.get(url).status_code == 403, url
        assert c.post("/users/save", data={"username": "x"}).status_code == 403
        assert c.post("/live/1/disconnect").status_code == 403
    with TestClient(app) as c:  # department manager: Finance in every office
        _login(c, "mgr_fin", "Fin123456")
        t = c.get("/users").text
        assert "yazd1" in t and "teh1" in t and "ali" not in t


def test_mikrotik_local_accounts_and_hotspot_deploy(env, tmp_path, monkeypatch):
    import io
    import threading
    import zipfile

    from pyftpdlib.authorizers import DummyAuthorizer
    from pyftpdlib.handlers import FTPHandler
    from pyftpdlib.servers import FTPServer

    from app import routeros
    from app.db import DeployLog, HotspotPackage
    from app.web import app
    from tests.fake_ros import FakeRouter

    # FTP server playing the router's file system, with an existing hotspot folder
    root = tmp_path / "router"
    (root / "flash" / "hotspot" / "img").mkdir(parents=True)
    (root / "flash" / "hotspot" / "login.html").write_text("<html><body>old</body></html>")
    (root / "flash" / "hotspot" / "img" / "old.png").write_bytes(b"PNG")
    auth = DummyAuthorizer()
    auth.add_user("acc", "pw", str(root), perm="elradfmw")
    handler = type("H", (FTPHandler,), {"authorizer": auth})
    srv = FTPServer(("127.0.0.1", 0), handler)
    port = srv.socket.getsockname()[1]
    threading.Thread(target=srv.serve_forever, kwargs={"timeout": 0.2}, daemon=True).start()

    fake = FakeRouter(ftp_root=str(root))
    fake.store["ip/hotspot/user/profile"] = [{".id": "*1", "name": "default", "shared-users": 1}]
    fake.store["ip/hotspot"] = [{".id": "*1", "name": "hotspot1"}]
    fake.store["ppp/profile"] = [{".id": "*1", "name": "default-encryption"}]
    fake.store["ip/hotspot/profile"] = [{".id": "*1", "name": "hsprof1", "html-directory": "flash/hotspot"}]
    monkeypatch.setattr(routeros, "api", fake.conn)

    db = SessionLocal()
    nas = db.query(Nas).filter_by(name="hq").one()
    nas.api_user, nas.api_password_enc, nas.api_host, nas.file_port = "acc", encrypt("pw"), "127.0.0.1", port
    db.commit()
    nid = nas.id
    db.close()
    try:
        with TestClient(app) as c:
            _login(c, "admin", "Secret12345")
            r = c.get(f"/mt-local?nas_id={nid}")
            assert r.status_code == 200 and "hotspot1" in r.text
            # create, then edit a hotspot user (empty password = keep, empty quota = remove)
            c.post("/mt-local/save", data={"nas_id": nid, "kind": "hs_user", "f_name": "guest1", "f_password": "abc123",
                                           "f_profile": "default", "f_server": "all", "f_limit-bytes-total": "1G",
                                           "f_limit-uptime": "10h", "f_comment": "مهمان"})
            u = fake.store["ip/hotspot/user"][0]
            assert u["limit-bytes-total"] == str(1024 ** 3) and u["password"] == "abc123" and u["comment"] == "مهمان"
            c.post("/mt-local/save", data={"nas_id": nid, "kind": "hs_user", "item_id": u[".id"], "f_name": "guest1",
                                           "f_password": "", "f_profile": "default", "f_server": "all",
                                           "f_limit-bytes-total": "", "f_limit-uptime": "20h", "f_comment": "مهمان"})
            u = fake.store["ip/hotspot/user"][0]
            assert u["password"] == "abc123" and "limit-bytes-total" not in u and u["limit-uptime"] == "20h"
            db = SessionLocal()
            assert db.query(User).filter_by(username="guest1").one().source == "router"
            db.close()
            # voucher-style bulk create
            r = c.post("/mt-local/bulk", data={"nas_id": nid, "kind": "hs_user", "mode": "generate", "prefix": "card",
                                               "start": "1", "count": "5", "pass_len": "6", "digits": "1",
                                               "f_profile": "default", "f_server": "all", "f_limit-bytes-total": "2G"})
            assert r.status_code == 200 and "card1" in r.text and "card5" in r.text
            assert len(fake.store["ip/hotspot/user"]) == 6
            ids = [i[".id"] for i in fake.store["ip/hotspot/user"] if i["name"].startswith("card")]
            c.post("/mt-local/action", data={"nas_id": nid, "kind": "hs_user", "action": "disable", "ids": ids[:2]})
            assert sum(1 for i in fake.store["ip/hotspot/user"] if i["disabled"]) == 2
            c.post("/mt-local/action", data={"nas_id": nid, "kind": "hs_user", "action": "delete", "ids": ids})
            assert [i["name"] for i in fake.store["ip/hotspot/user"]] == ["guest1"]
            # VPN secret + profile
            c.post("/mt-local/save", data={"nas_id": nid, "kind": "ppp_secret", "f_name": "vpnuser", "f_password": "Vp1",
                                           "f_service": "l2tp", "f_profile": "default-encryption",
                                           "f_remote-address": "10.10.10.5"})
            assert fake.store["ppp/secret"][0]["service"] == "l2tp"
            c.post("/mt-local/save", data={"nas_id": nid, "kind": "ppp_profile", "f_name": "vpn-4M",
                                           "f_rate-limit": "2M/4M", "f_only-one": "yes"})
            assert fake.store["ppp/profile"][-1]["rate-limit"] == "2M/4M"
            assert "ایجاد اکانت‌های VPN" in c.get("/audit").text

            # hotspot folder: upload a zip whose files sit in a top "hotspot/" folder
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w") as zf:
                zf.writestr("hotspot/login.html", "<html><body>NEW $(link-login-only)</body></html>")
                zf.writestr("hotspot/img/logo.png", b"\x89PNG")
                zf.writestr("__MACOSX/hotspot/._login.html", b"junk")
            r = c.post("/hotspot/upload", data={"name": "new-login", "inject": "1", "public_url": "http://10.0.0.5:8080"},
                       files={"file": ("hs.zip", buf.getvalue(), "application/zip")})
            assert r.status_code == 200
            db = SessionLocal()
            pkg = db.query(HotspotPackage).filter_by(name="new-login").one()
            assert pkg.files == 2
            db.close()
            r = c.post("/hotspot/deploy", data={"package_id": pkg.id, "nas_ids": nid, "backup": "1"})
            assert r.status_code == 200
            login = (root / "flash" / "hotspot" / "login.html").read_text()
            assert "NEW" in login and "http://10.0.0.5:8080/ua?mac=$(mac-esc)" in login
            assert (root / "flash" / "hotspot" / "img" / "logo.png").read_bytes() == b"\x89PNG"
            assert (root / "flash" / "hotspot" / "img" / "old.png").exists()       # no mirror -> kept
            db = SessionLocal()
            rec = db.query(DeployLog).order_by(DeployLog.id.desc()).first()
            assert rec.ok and rec.target_dir == "flash/hotspot" and rec.backup_id
            backup = db.get(HotspotPackage, rec.backup_id)
            db.close()
            zf = zipfile.ZipFile(io.BytesIO(c.get(f"/hotspot/{backup.id}/download").content))
            assert zf.read("login.html") == b"<html><body>old</body></html>" and "img/old.png" in zf.namelist()
            # edit a file in the package, then roll back with mirror (old.png restored, logo.png removed)
            c.post(f"/hotspot/{pkg.id}/edit", data={"path": "status.html", "content": "<p>ok</p>"})
            assert "status.html" in c.get(f"/hotspot/{pkg.id}/files").text
            c.post("/hotspot/deploy", data={"package_id": backup.id, "nas_ids": nid, "mirror": "1"})
            assert (root / "flash" / "hotspot" / "login.html").read_text() == "<html><body>old</body></html>"
            assert not (root / "flash" / "hotspot" / "img" / "logo.png").exists()
            assert "موفق" in c.get("/hotspot").text
        with TestClient(app) as c:  # IT expert may use both features, HR may not
            _login(c, "it", "ItPass2026x")
            assert c.get(f"/mt-local?nas_id={nid}").status_code == 200
            assert c.get("/hotspot").status_code == 200
        with TestClient(app) as c:
            _login(c, "hr", "HrPass2026x")
            assert c.get("/mt-local").status_code == 403 and c.get("/hotspot").status_code == 403
    finally:
        srv.close_all()


def test_accounting_creates_router_local_user(env):
    req = env.CreateAcctPacket(User_Name="mt_local_only")
    req["Acct-Status-Type"] = 1
    req["Acct-Session-Id"] = "L1"
    req["Framed-IP-Address"] = "10.5.50.77"
    assert env.SendPacket(req).code == packet.AccountingResponse
    db = SessionLocal()
    u = db.query(User).filter_by(username="mt_local_only").one()
    assert u.source == "router"
    assert db.query(AcctSession).filter_by(acct_session_id="L1").one().user_id == u.id
    db.close()
    auth = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="mt_local_only")
    auth["User-Password"] = auth.PwCrypt("")
    assert env.SendPacket(auth).code == packet.AccessReject


def test_internal_vs_internet_split(env, monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    from app import routeros
    from app.reports import split_usage
    from app.web import app
    from tests.fake_ros import FakeRouter

    fake = FakeRouter()
    fake.store["interface"] = [{".id": "*1", "name": "ether1", "type": "ether", "running": True},
                               {".id": "*7", "name": "bridge-lan", "type": "bridge", "running": True},
                               {".id": "*8", "name": "ether3", "type": "ether", "running": True}]
    fake.store["ip/address"] = [{"address": "5.5.5.5/29", "interface": "ether1", "dynamic": True},
                                {"address": "10.5.50.1/24", "interface": "bridge-lan"},
                                {"address": "192.168.100.1/24", "interface": "ether3"}]
    fake.store["ip/route"] = [{"dst-address": "0.0.0.0/0", "gateway": "5.5.5.1", "active": True}]
    fake.store["ip/pool"] = [{"name": "vpn", "ranges": "10.10.10.2-10.10.10.254"}]
    monkeypatch.setattr(routeros, "api", fake.conn)
    db = SessionLocal()
    nas = db.query(Nas).filter_by(name="hq").one()
    nid = nas.id
    db.close()

    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        assert "تعیین نشده" in c.get("/nas").text
        c.post(f"/nas/{nid}/detect-net")
        r = c.get(f"/nas/{nid}/edit")
        assert 'value="ether1" checked' in r.text and 'value="bridge-lan" >' in r.text
        assert "10.5.50.0/24" in r.text and "10.10.10.2-10.10.10.254" in r.text
        r = c.post("/nas/save", data={"id": nid, "name": "hq", "ip": "127.0.0.1", "wan": "ether1",
                                      "lan_ranges": "192.168.100.0/24\n203.0.113.0/24\nnot-a-range"},
                   follow_redirects=False)
        assert r.status_code == 303
        assert "not-a-range" in c.get("/nas").text            # invalid range reported
        c.post("/nas/scope", data={"internal_ranges": "10.0.0.0/8\n192.168.0.0/16", "internet_ranges": "10.99.0.5"})
        db = SessionLocal()
        nas = db.get(Nas, nid)
        assert nas.wan_interfaces == "ether1" and nas.lan_ranges == "192.168.100.0/24\n203.0.113.0/24"
        db.close()

        _acct(env, 1, "S9", 0, 0, 0)
        time.sleep(1.1)
        _acct(env, 3, "S9", 3_000_000, 20_000_000, 60)

        fields = [(1, 8), (4, 1), (7, 2), (8, 4), (11, 2), (12, 4), (10, 2), (14, 2)]
        tpl = struct.pack("!HH", 300, len(fields)) + b"".join(struct.pack("!HH", t, ln) for t, ln in fields)

        def rec(src, dst, octets, in_if, out_if, sport=445, dport=50000):
            return (struct.pack("!QBH", octets, 6, sport) + socket.inet_aton(src) + struct.pack("!H", dport) +
                    socket.inet_aton(dst) + struct.pack("!HH", in_if, out_if))
        MB = 1024 * 1024
        recs = [rec("192.168.100.20", "10.5.50.10", 5 * MB, 8, 7),      # file server -> user: internal
                rec("10.5.50.10", "192.168.100.20", 2 * MB, 7, 8),      # user -> file server: internal
                rec("203.0.113.9", "10.5.50.10", 1 * MB, 1, 7),         # router's own internal range: internal
                rec("1.1.1.1", "10.5.50.10", 4 * MB, 1, 7, 443),        # internet via WAN
                rec("10.99.0.5", "10.5.50.10", 3 * MB, 8, 7, 8080),     # proxy exception: internet
                rec("192.168.1.1", "10.5.50.10", 1 * MB, 1, 7, 80)]     # private but via WAN: internet
        body = b"".join(recs)
        pkt = lambda sets: struct.pack("!HHIII", 10, 16 + len(sets), int(time.time()), 1, 0) + sets  # noqa: E731
        col = Collector()
        col.handle(pkt(struct.pack("!HH", 2, 4 + len(tpl)) + tpl), "127.0.0.1")
        col.handle(pkt(struct.pack("!HH", 300, 4 + len(body)) + body), "127.0.0.1")
        col.flush()

        db = SessionLocal()
        ali = db.query(User).filter_by(username="ali").one()
        lan = db.query(TrafficCategory).filter_by(user_id=ali.id, category="lan").one()
        assert (lan.bytes_down, lan.bytes_up) == (6 * MB, 2 * MB)
        from datetime import datetime, timedelta

        from app.db import now
        today = datetime.combine(now().date(), datetime.min.time())
        sp = split_usage(db, today, today + timedelta(days=1), [ali.id])
        acct_total = sum((s.d_in or 0) + (s.d_out or 0) for s in db.query(SpeedSample).filter_by(user_id=ali.id))
        assert sp["internal"] == 8 * MB and sp["internet"] == acct_total - 8 * MB
        assert sp["user_daily"][0]["internal_down"] == 6 * MB
        db.close()

        r = c.get("/reports?preset=today")
        assert "مصرف روزانه: اینترنت / تانل / داخل شبکه" in r.text and "شبکه داخلی" in r.text
        r = c.get("/reports/export?kind=daily&preset=today")
        assert "داخل شبکه (بایت)" in r.text and f",{8 * MB}," in r.text
        assert "مصرف داخل شبکه" in c.get("/mgmt?preset=today").text
        wb = load_workbook(BytesIO(c.get("/mgmt/export.xlsx?preset=today").content))
        ws = wb["مصرف روزانه کاربران"]
        row = [cell.value for cell in ws[2]]
        assert row[1] == "ali" and row[6] == round(8 * MB / 1024 ** 3, 3)
        assert "داخل شبکه" in c.get("/attendance?preset=today").text
        wb = load_workbook(BytesIO(c.get("/attendance/export.xlsx?preset=today").content))
        assert "داخل شبکه (GB)" in [cell.value for cell in wb["ساعات اتصال"][3]]
        assert "داخل شبکه" in c.get("/").text


def _mock_ad(entries):
    from ldap3 import MOCK_SYNC, Connection, Server
    c = Connection(Server("dc1"), user="CN=svc,CN=Users,DC=corp,DC=local", password="pw", client_strategy=MOCK_SYNC)
    c.strategy.add_entry("CN=svc,CN=Users,DC=corp,DC=local", {"objectClass": ["user"], "userPassword": "pw"})
    for dn, attrs in entries:
        c.strategy.add_entry(dn, attrs)
    c.bind()
    return c


def test_active_directory_base_dn_autodetect_and_groups(env, monkeypatch):
    from app import ldap_ad
    from app.db import get_setting, set_setting
    from app.web import app
    person = {"objectClass": ["top", "person", "user"], "objectCategory": "person", "userAccountControl": 512}
    entries = [("CN=Sara Ahmadi,OU=Staff,DC=corp,DC=local",
                {**person, "sAMAccountName": "sara.ad", "displayName": "سارا احمدی", "department": "مالی",
                 "physicalDeliveryOfficeName": "دفتر یزد", "memberOf": ["CN=Net-Users,OU=Groups,DC=corp,DC=local"]}),
               ("CN=Old,OU=Staff,DC=corp,DC=local", {**person, "sAMAccountName": "old.ad", "userAccountControl": 514}),
               ("CN=Net-Users,OU=Groups,DC=corp,DC=local",
                {"objectClass": ["top", "group"], "cn": "Net-Users",
                 "member": ["CN=Sara Ahmadi,OU=Staff,DC=corp,DC=local"]})]
    monkeypatch.setattr(ldap_ad, "_connect", lambda c, timeout=30: _mock_ad(entries))
    db = SessionLocal()
    for k, v in {"ldap_server": "dc1", "ldap_domain": "corp.local", "ldap_bind_user": "svc", "ldap_base_dn": "",
                 "ldap_user_filter": ""}.items():
        set_setting(db, k, v)
    db.commit()
    ok, msg = ldap_ad.test_connection(db)
    assert ok and "DC=corp,DC=local" in msg and "2 کاربر" in msg, msg
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.post("/ad/groups")
        assert r.status_code == 200 and "Net-Users" in r.text and "CN=Net-Users,OU=Groups,DC=corp,DC=local" in r.text
        assert len(r.headers.get("set-cookie", "")) < 3000                 # nothing big put in the session cookie
        r = c.post("/ad/sync")
        assert "2 کاربر در AD" in r.text
    sara = db.query(User).filter_by(username="sara.ad").one()
    assert sara.source == "ldap" and sara.full_name == "سارا احمدی" and sara.office.name == "دفتر یزد"
    assert db.query(User).filter_by(username="old.ad").one().enabled is False
    # an empty result (wrong Base DN) must not disable every AD user
    monkeypatch.setattr(ldap_ad, "_connect", lambda c, timeout=30: _mock_ad([]))
    set_setting(db, "ldap_base_dn", "OU=Nowhere,DC=corp,DC=local")
    db.commit()
    with pytest.raises(RuntimeError):
        ldap_ad.sync_users(db)
    db.expire_all()
    assert db.query(User).filter_by(username="sara.ad").one().enabled is True
    set_setting(db, "ldap_server", "")
    db.commit()
    assert get_setting(db, "ldap_server") == ""
    db.close()


def test_user_self_service_portal(env):
    from app import web
    from app.db import decrypt
    from app.web import app
    with TestClient(app) as c:
        assert "مصرف اینترنت من" in c.get("/my").text
        r = c.post("/my/login", data={"username": "ali", "password": "wrong"})
        assert "نام کاربری یا رمز عبور اشتباه است" in r.text
        r = c.post("/my/login", data={"username": "ALI", "password": "Pass123"})
        assert "سلام" in r.text and "حجم باقی‌مانده" in r.text and "اتصال‌های اخیر" in r.text
        assert "مصرف روزانه (۳۰ روز اخیر)" in r.text
        # the portal session gives no access to the admin panel
        assert c.get("/users", follow_redirects=False).headers["location"] == "/login"
        c.post("/my/login", data={"username": "ali", "password": "Pass123"})
        r = c.post("/my/password", data={"current": "Pass123", "new": "abc", "again": "abc"})
        assert "حداقل ۶" in r.text
        r = c.post("/my/password", data={"current": "bad", "new": "NewPass99", "again": "NewPass99"})
        assert "رمز فعلی درست نیست" in r.text
        r = c.post("/my/password", data={"current": "Pass123", "new": "NewPass99", "again": "NewPass99"})
        assert "رمز عبور عوض شد" in r.text
        db = SessionLocal()
        assert decrypt(db.query(User).filter_by(username="ali").one().password_enc) == "NewPass99"
        db.close()
        req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="ali")
        req["User-Password"] = req.PwCrypt("NewPass99")
        assert env.SendPacket(req).code == packet.AccessAccept
        # router-only accounts have no password here
        assert "اشتباه" in c.post("/my/login", data={"username": "mt_local_only", "password": ""}).text
        c.get("/my/logout")
        assert "نام کاربری" in c.get("/my").text
        web.MY_FAILS.clear()
        for _ in range(5):
            c.post("/my/login", data={"username": "ali", "password": "x"})
        r = c.post("/my/login", data={"username": "ali", "password": "NewPass99"})
        assert "۱۰ دقیقه" in r.text and "سلام" not in r.text
        web.MY_FAILS.clear()
    db = SessionLocal()
    db.query(User).filter_by(username="ali").one().password_enc = encrypt("Pass123")
    db.commit()
    db.close()


def test_router_api_sessions_and_device_history(env, monkeypatch):
    """Local hotspot / PPP logins without RADIUS accounting are tracked from the router's active lists,
    and every device keeps its IP / user history with time ranges."""
    from datetime import timedelta

    from app import devices, jalali, jobs, routeros
    from app.db import Device, DeviceHistory, now
    from app.web import app
    from tests.fake_ros import FakeRouter

    fake = FakeRouter()
    fake.store["system/identity"] = [{"name": "Branch"}]
    fake.store["ip/hotspot/active"] = [{".id": "*A1", "user": "guest1", "address": "10.5.50.20",
                                        "mac-address": "aa:bb:cc:00:00:01", "uptime": "5m", "bytes-in": "1000",
                                        "bytes-out": "50000"}]
    fake.store["ppp/active"] = [{".id": "*B1", "name": "vpnuser", "service": "l2tp", "caller-id": "",
                                 "address": "10.10.10.5", "uptime": "1h2m3s"}]
    fake.store["interface"] = [{".id": "*20", "name": "<l2tp-vpnuser>", "rx-byte": "2000", "tx-byte": "90000"}]
    fake.store["ip/arp"] = [{"address": "10.5.50.30", "mac-address": "AA:BB:CC:00:00:02", "status": "reachable"},
                            {"address": "10.5.50.31", "mac-address": "AA:BB:CC:00:00:03", "status": "failed"}]
    monkeypatch.setattr(routeros, "api", fake.conn)
    db = SessionLocal()
    nas = Nas(name="branch", ip="0.0.0.0", secret="s", api_host="10.9.9.1", api_user="api")
    db.add(nas)
    db.commit()
    nid = nas.id
    db.close()

    jobs.poll_routers()
    db = SessionLocal()
    hs = db.query(AcctSession).filter_by(username="guest1", active=True).one()
    assert hs.acct_session_id.startswith(f"api:{nid}:hs:") and hs.nas_ip == "10.9.9.1" and hs.service == "hotspot"
    assert hs.bytes_out == 50000 and hs.mac == "AA:BB:CC:00:00:01"
    assert (now() - hs.start_time).total_seconds() >= 299
    assert db.query(User).filter_by(username="guest1").one().source == "router"
    vpn = db.query(AcctSession).filter_by(username="vpnuser", active=True).one()
    assert vpn.service == "vpn" and vpn.bytes_in == 2000 and vpn.bytes_out == 90000
    d1 = db.query(Device).filter_by(mac="AA:BB:CC:00:00:01").one()
    assert d1.last_ip == "10.5.50.20" and d1.last_user == "guest1" and d1.last_nas_id == nid
    assert db.query(Device).filter_by(mac="AA:BB:CC:00:00:02").one().last_ip == "10.5.50.30"
    assert db.query(Device).filter_by(mac="AA:BB:CC:00:00:03").first() is None      # failed ARP entry ignored
    assert "1 هات‌اسپات · 1 VPN" in jobs.live["routers"]["branch"]
    db.close()

    # second poll: counters grew -> speed sample; the hotspot user moved to a new IP
    fake.store["ip/hotspot/active"][0].update({"bytes-in": "3000", "bytes-out": "150000", "uptime": "6m",
                                               "address": "10.5.50.21"})
    jobs.poll_routers()
    db = SessionLocal()
    hs = db.query(AcctSession).filter_by(username="guest1", active=True).one()
    assert hs.bytes_out == 150000
    assert db.query(SpeedSample).filter_by(session_id=hs.id).count() == 2
    ips = {h.value for h in db.query(DeviceHistory).filter_by(mac="AA:BB:CC:00:00:01", kind="ip")}
    assert ips == {"10.5.50.20", "10.5.50.21"}
    users = db.query(DeviceHistory).filter_by(mac="AA:BB:CC:00:00:01", kind="user").all()
    assert len(users) == 1 and users[0].value == "guest1" and users[0].nas_id == nid
    db.close()

    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/live")
        assert "guest1" in r.text and "vpnuser" in r.text and "branch" in r.text and ">API<" in r.text
        r = c.get("/")
        assert "guest1" in r.text or "کاربران آنلاین" in r.text
        r = c.get("/devices", params={"q": "10.5.50.20"})
        assert "AA:BB:CC:00:00:01" in r.text and "سابقه منطبق" in r.text
        r = c.get("/devices", params={"q": "guest1"})
        assert "AA:BB:CC:00:00:01" in r.text
        r = c.get("/devices/aa-bb-cc-00-00-01")
        assert r.status_code == 200 and "10.5.50.21" in r.text and "guest1" in r.text and "هم‌اکنون" in r.text
        r = c.get("/devices", params={"at": jalali.to_jalali_str(now(), True)})
        assert "AA:BB:CC:00:00:01" in r.text
        assert "AA:BB:CC:00:00:01" in c.get(f"/users/{hs.user_id}/edit").text
        r = c.get("/nas")
        assert "عیب‌یابی" in r.text and "127.0.0.1" in r.text and "1 / 2" not in r.text

    # the login disappears from the router -> session closed; the device history stays
    fake.store["ip/hotspot/active"] = []
    jobs.poll_routers()
    db = SessionLocal()
    s = db.get(AcctSession, hs.id)
    assert s.active is False and s.terminate_cause == "Logout"
    # RADIUS accounting already covering a login -> no API duplicate
    fake.store["ip/hotspot/active"] = [{".id": "*A2", "user": "ali", "address": "10.5.50.40",
                                        "mac-address": "AA:BB:CC:00:00:04", "uptime": "1m", "bytes-in": "1",
                                        "bytes-out": "1"}]
    db.add(AcctSession(acct_session_id="rad-1", nas_ip="10.9.9.1", username="ali", framed_ip="10.5.50.40",
                       active=True, start_time=now(), last_update=now()))
    db.commit()
    db.close()
    jobs.poll_routers()
    db = SessionLocal()
    assert db.query(AcctSession).filter(AcctSession.username == "ali", AcctSession.active.is_(True),
                                        AcctSession.acct_session_id.like("api:%")).count() == 0
    db.query(AcctSession).filter_by(acct_session_id="rad-1").one().active = False

    # history intervals: a gap longer than DEVICE_HISTORY_GAP_MIN starts a new interval
    t0 = now() - timedelta(hours=3)
    devices._open.clear()
    for minutes in (0, 5, 10, 60, 65):
        devices.record(db, "AA:BB:CC:00:00:09", ip="10.5.50.99", nas_id=nid, source="arp",
                       ts=t0 + timedelta(minutes=minutes))
    rows = db.query(DeviceHistory).filter_by(mac="AA:BB:CC:00:00:09").order_by(DeviceHistory.first_seen).all()
    assert [(int((r.first_seen - t0).total_seconds() // 60), int((r.last_seen - t0).total_seconds() // 60))
            for r in rows] == [(0, 10), (60, 65)]
    # devices never disappear; not seen for longer than device_stale_days -> stale
    d = db.query(Device).filter_by(mac="AA:BB:CC:00:00:09").one()
    d.last_seen = now() - timedelta(days=40)
    db.commit()
    db.close()
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/devices", params={"status": "stale"})
        assert "AA:BB:CC:00:00:09" in r.text and "AA:BB:CC:00:00:01" not in r.text
        r = c.post("/devices/settings", data={"device_stale_days": "60"}, follow_redirects=True)
        assert "AA:BB:CC:00:00:09" not in c.get("/devices", params={"status": "stale"}).text
    db = SessionLocal()
    db.delete(db.get(Nas, nid))
    db.commit()
    db.close()


def test_bulk_edit_and_mobile_written_to_active_directory(env, monkeypatch):
    from app import ldap_ad
    from app.db import Office, set_setting
    from app.web import app
    person = {"objectClass": ["top", "person", "user"], "objectCategory": "person", "userAccountControl": 512}
    basic, vip = "CN=Net-Basic,OU=Groups,DC=corp,DC=local", "CN=Net-VIP,OU=Groups,DC=corp,DC=local"
    ali_dn, reza_dn, nima_dn = ("CN=Ali K,OU=Staff,DC=corp,DC=local", "CN=Reza M,OU=Staff,DC=corp,DC=local",
                                "CN=Nima R,OU=Staff,DC=corp,DC=local")
    entries = [(ali_dn, {**person, "sAMAccountName": "ali.k", "displayName": "Ali K", "mobile": "09120000000",
                         "title": "Developer", "memberOf": [basic]}),
               (reza_dn, {**person, "sAMAccountName": "reza.m", "displayName": "Reza M"}),
               (nima_dn, {**person, "sAMAccountName": "nima.r", "displayName": "Nima R"}),
               (basic, {"objectClass": ["top", "group"], "cn": "Net-Basic", "member": [ali_dn]}),
               (vip, {"objectClass": ["top", "group"], "cn": "Net-VIP"})]
    conn = _mock_ad(entries)
    conn.unbind = lambda *a, **k: True
    real_modify = conn.modify

    def modify(dn, changes, *a, **k):           # the service account may not write to Nima's account
        if dn == nima_dn:
            conn.result = {"result": 50, "description": "insufficientAccessRights", "message": ""}
            return False
        return real_modify(dn, changes, *a, **k)
    conn.modify = modify
    monkeypatch.setattr(ldap_ad, "_connect", lambda c, timeout=30: conn)
    db = SessionLocal()
    for k, v in {"ldap_server": "dc1", "ldap_domain": "corp.local", "ldap_bind_user": "svc",
                 "ldap_base_dn": "DC=corp,DC=local", "ldap_user_filter": "(sAMAccountName=*.*)",
                 "ldap_default_group": ""}.items():
        set_setting(db, k, v)
    g_basic = Group(name="AD-Basic", ldap_group_dn=basic)
    g_vip = Group(name="AD-VIP", ldap_group_dn=vip)
    office = Office(name="دفتر ونک")
    db.add_all([g_basic, g_vip, office])
    db.commit()
    gid_vip, oid = g_vip.id, office.id
    ldap_ad.sync_users(db)
    ali = db.query(User).filter_by(username="ali.k").one()
    assert ali.mobile == "09120000000" and ali.title == "Developer" and ali.group.name == "AD-Basic"
    ids = [db.query(User).filter_by(username=n).one().id for n in ("ali.k", "reza.m", "nima.r", "night")]
    db.close()

    def ad(dn, attr):
        conn.search(dn, "(objectClass=*)", "BASE", attributes=[attr])
        vals = conn.entries[0][attr].values if conn.entries and attr in conn.entries[0].entry_attributes else []
        return [str(v) for v in vals]

    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/users")
        assert 'id="bulkedit"' in r.text and "09120000000" in r.text
        r = c.post("/users/bulk-edit", data={
            "ids": [str(i) for i in ids], "set": ["group_id", "office_id", "mobile", "city", "title"],
            "v_group_id": str(gid_vip), "v_office_id": str(oid), "v_mobile": "۰۹۱۲۱۱۱۲۲۳۳", "v_city": "Tehran",
            "v_title": "", "ad_write": "1", "ad_exclusive": "1", "back": "/users"})
        assert r.status_code == 200 and "3</b> از 4" in r.text and "Delegate Control" in r.text
        assert ad(ali_dn, "mobile") == ["09121112233"] and ad(ali_dn, "l") == ["Tehran"]
        assert ad(ali_dn, "physicalDeliveryOfficeName") == ["دفتر ونک"] and ad(ali_dn, "title") == []
        assert set(ad(vip, "member")) == {ali_dn, reza_dn} and ali_dn not in ad(basic, "member")
        db = SessionLocal()
        ali = db.query(User).filter_by(username="ali.k").one()
        assert ali.group_id == gid_vip and ali.office_id == oid and ali.mobile == "09121112233" and ali.title == ""
        nima = db.query(User).filter_by(username="nima.r").one()
        assert nima.group_id is None and nima.mobile == ""                  # failed in AD -> nothing changed
        night = db.query(User).filter_by(username="night").one()
        assert night.mobile == "09121112233" and night.city == "Tehran"    # local user: only here
        db.close()
        # single user form: mobile written back to AD
        r = c.post("/users/save", data={"id": str(ids[1]), "username": "reza.m", "full_name": "Reza M",
                                        "mobile": "09350000000", "group_id": str(gid_vip), "enabled": "1",
                                        "ad_write": "1"}, follow_redirects=True)
        assert "Active Directory هم ثبت شد" in r.text
        assert ad(reza_dn, "mobile") == ["09350000000"]
        r = c.post("/users/save", data={"id": str(ids[2]), "username": "nima.r", "mobile": "0911",
                                        "enabled": "1", "ad_write": "1"}, follow_redirects=True)
        assert "ناموفق" in r.text
    db = SessionLocal()
    assert db.query(User).filter_by(username="nima.r").one().mobile == ""
    # the next AD sync keeps the values written to AD
    ldap_ad.sync_users(db)
    ali = db.query(User).filter_by(username="ali.k").one()
    assert ali.mobile == "09121112233" and ali.city == "Tehran" and ali.office_id == oid
    for k in ("ldap_server", "ldap_base_dn", "ldap_user_filter"):
        set_setting(db, k, "")
    db.commit()
    db.close()


def test_review_fixes(env, monkeypatch):
    from app import auth, jobs, ldap_ad, routeros
    from app.db import AuthLog, PanelUser, get_setting, now, set_setting
    from app.web import app
    from tests.fake_ros import FakeRouter

    # numeric hotspot user names / DHCP host names (librouteros returns them as int)
    fake = FakeRouter()
    fake.store["ip/hotspot/active"] = [{".id": "*C1", "user": "1001", "address": "10.7.0.5",
                                        "mac-address": "AA:BB:CC:00:10:01", "uptime": "2m", "bytes-in": "10",
                                        "bytes-out": "20"}]
    fake.store["ip/dhcp-server/lease"] = [{"mac-address": "AA:BB:CC:00:10:02", "address": "10.7.0.6",
                                           "host-name": "100234", "status": "bound", "last-seen": "10s"}]
    monkeypatch.setattr(routeros, "api", fake.conn)
    db = SessionLocal()
    nas = Nas(name="numeric", ip="10.7.0.1", secret="s", api_user="api")
    db.add(nas)
    db.commit()
    nid = nas.id
    db.close()
    jobs.poll_routers()
    db = SessionLocal()
    assert db.query(AcctSession).filter_by(username="1001", active=True).count() == 1
    assert "OK" in jobs.live["routers"]["numeric"]
    from app.db import Device
    assert db.query(Device).filter_by(mac="AA:BB:CC:00:10:02").one().hostname == "100234"
    db.delete(db.get(Nas, nid))
    db.query(AcctSession).filter_by(username="1001").update({"active": False})
    db.commit()

    # '%' / '_' in a RADIUS User-Name are not wildcards
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="%")
    req["User-Password"] = req.PwCrypt("Pass123")
    assert env.SendPacket(req).code == packet.AccessReject
    req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name="al_")
    req["User-Password"] = req.PwCrypt("Pass123")
    assert env.SendPacket(req).code == packet.AccessReject

    # AD: an account outside the configured filter is not let in on its first login
    person = {"objectClass": ["top", "person", "user"], "objectCategory": "person", "userAccountControl": 512}
    entries = [("CN=In,OU=Staff,DC=corp,DC=local", {**person, "sAMAccountName": "in.user", "department": "IT"}),
               ("CN=Out,OU=Staff,DC=corp,DC=local", {**person, "sAMAccountName": "out.user", "department": "HR"})]
    monkeypatch.setattr(ldap_ad, "_connect", lambda c, timeout=30: _mock_ad(entries))
    monkeypatch.setattr(ldap_ad, "authenticate", lambda db, u, p: (True, ""))
    for k, v in {"ldap_server": "dc1", "ldap_domain": "corp.local", "ldap_bind_user": "svc",
                 "ldap_base_dn": "DC=corp,DC=local", "ldap_user_filter": "(department=IT)"}.items():
        set_setting(db, k, v)
    db.commit()
    for name, code in (("out.user", packet.AccessReject), ("IN.user", packet.AccessAccept)):
        req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name=name)
        req["User-Password"] = req.PwCrypt("whatever")
        assert env.SendPacket(req).code == code, name
    db.expire_all()
    assert db.query(User).filter_by(username="out.user").first() is None
    assert db.query(User).filter_by(username="in.user").one().source == "ldap"
    assert "Base DN" in db.query(AuthLog).filter_by(username="out.user").order_by(AuthLog.id.desc()).first().reason
    for k in ("ldap_server", "ldap_base_dn", "ldap_user_filter"):
        set_setting(db, k, "")
    db.commit()

    # old installs: it / hr with the well-known password get disabled once
    it = db.query(PanelUser).filter_by(username="it").one()
    old_hash, old_enabled, old_must = it.password_hash, it.enabled, it.must_change_password
    it.password_hash, it.enabled, it.must_change_password = auth.hash_password("it12345678"), True, True
    set_setting(db, "default_accounts_secured", "")
    db.commit()
    auth.secure_default_accounts(db)
    assert db.query(PanelUser).filter_by(username="it").one().enabled is False
    assert get_setting(db, "default_accounts_secured") == "1"
    it = db.query(PanelUser).filter_by(username="it").one()
    it.password_hash, it.enabled, it.must_change_password = old_hash, old_enabled, old_must
    db.commit()

    # a deleted user's usage does not move to the next user that gets the same id
    u = User(username="tmp.del", password_enc="")
    db.add(u)
    db.commit()
    uid = u.id
    db.add(SpeedSample(session_id=0, user_id=uid, ts=now(), seconds=60, d_in=0, d_out=900 * 1024 * 1024))
    db.commit()
    db.close()
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        c.post(f"/users/{uid}/delete")
        db = SessionLocal()
        assert db.query(SpeedSample).filter_by(user_id=uid).count() == 0
        db.close()
        # local redirects only; bad input gives a message, not an error page
        r = c.post("/settings", data={"work_hours": "sat-wed 08:00-16:00", "back": "https://evil.example/"},
                   follow_redirects=False)
        assert r.headers["location"] == "/mgmt"
        assert c.get("/reports", params={"from": "1404/13/45", "to": "x"}).status_code == 200
        r = c.post("/nas/save", data={"name": "dup", "ip": "127.0.0.1", "secret": "x"}, follow_redirects=True)
        assert r.status_code == 200 and "با همین IP" in r.text
        r = c.post("/users/save", data={"username": "baddate", "expire_at": "1404/99/99", "enabled": "1"},
                   follow_redirects=True)
        assert r.status_code == 200 and "تاریخ انقضا" in r.text


def test_radius_session_kept_up_to_date_from_router_api(env, monkeypatch):
    """Router sends RADIUS accounting but rare / no Interim-Updates: usage comes from the API counters every poll
    and the final Stop only adds the rest (no double counting)."""
    from sqlalchemy import func

    from app import jobs, routeros
    from tests.fake_ros import FakeRouter
    assert _acct(env, 1, "FEED1", 0, 0, 0).code == packet.AccountingResponse
    fake = FakeRouter()
    fake.store["ip/hotspot/active"] = [{".id": "*F1", "user": "ali", "address": "10.5.50.10",
                                        "mac-address": "3C:22:FB:11:22:33", "uptime": "40s", "bytes-in": "1000",
                                        "bytes-out": "5000000"}]
    # the RADIUS packets above came from 127.0.0.1: one of this router's addresses
    fake.store["ip/address"] = [{".id": "*A1", "address": "127.0.0.1/8", "interface": "lo"}]
    monkeypatch.setattr(routeros, "api", fake.conn)
    db = SessionLocal()
    nas = Nas(name="feed", ip="10.6.0.1", secret="s", api_user="api")
    db.add(nas)
    db.commit()
    nid = nas.id
    db.close()
    jobs.poll_routers()
    db = SessionLocal()
    s = db.query(AcctSession).filter_by(acct_session_id="FEED1").one()
    assert s.active and s.bytes_out == 5_000_000
    assert db.query(AcctSession).filter(AcctSession.acct_session_id.like("api:%"), AcctSession.username == "ali",
                                        AcctSession.active.is_(True)).count() == 0
    db.close()
    _acct(env, 2, "FEED1", 2000, 7_000_000, 60)
    db = SessionLocal()
    s = db.query(AcctSession).filter_by(acct_session_id="FEED1").one()
    total = db.query(func.sum(SpeedSample.d_out)).filter(SpeedSample.session_id == s.id).scalar()
    assert not s.active and s.bytes_out == 7_000_000 and total == 7_000_000
    db.delete(db.get(Nas, nid))
    db.commit()
    db.close()


def test_routers_behind_nat_told_apart(env, monkeypatch):
    """Offices behind NAT: RADIUS and Traffic-Flow of every router arrive from the same address (here
    127.0.0.1, which is also the main router). NAS-Identifier / NAS-IP-Address and the secret tell them apart."""
    from app import netscope
    from app.db import AuthLog
    from app.netflow import Collector
    from app.web import app
    db = SessionLocal()
    branch = Nas(name="tehran", ip="10.20.0.1", secret="branch-secret", identifier="Office-Tehran",
                 lan_ranges="100.100.0.0/16")
    db.add(branch)
    db.add(User(username="sara", password_enc=encrypt("Sara1234")))
    db.add(User(username="hq_staff", password_enc=encrypt("x")))
    db.commit()
    bid = branch.id
    hq_id = db.query(Nas).filter_by(name="hq").one().id
    db.close()
    radius_server.nas_cache.invalidate()
    netscope.scope.invalidate()
    bcl = Client(server="127.0.0.1", authport=18121, acctport=18131, secret=b"branch-secret",
                 dict=Dictionary(config.RADIUS_DICTIONARY))
    bcl.timeout, bcl.retries = 2, 1

    def auth(cl, ident=None, nas_ip=None):
        req = cl.CreateAuthPacket(code=packet.AccessRequest, User_Name="sara")
        req["User-Password"] = req.PwCrypt("Sara1234")
        if ident:
            req["NAS-Identifier"] = ident
        if nas_ip:
            req["NAS-IP-Address"] = nas_ip
        req.add_message_authenticator()
        return cl.SendPacket(req)
    assert auth(bcl, "Office-Tehran", "10.20.0.1").code == packet.AccessAccept
    assert auth(env, "hq-router").code == packet.AccessAccept           # main router still itself
    # no NAS-Identifier: the secret alone finds the branch
    assert auth(bcl).code == packet.AccessAccept
    db = SessionLocal()
    logs = [r.nas_ip for r in db.query(AuthLog).filter_by(username="sara").order_by(AuthLog.id)]
    assert logs[-3:] == ["10.20.0.1", "127.0.0.1", "10.20.0.1"]
    db.close()

    def acct(cl, status, sid, ip, ident="Office-Tehran", out=0):
        req = cl.CreateAcctPacket(User_Name="sara")
        req["Acct-Status-Type"] = status
        req["Acct-Session-Id"] = sid
        req["Framed-IP-Address"] = ip
        req["NAS-Identifier"] = ident
        req["Acct-Output-Octets"] = out
        return cl.SendPacket(req)
    assert acct(bcl, 1, "NAT1", "192.168.88.20").code == packet.AccountingResponse
    # a router that is not defined (unknown identity, unknown secret): dropped, shown in diagnostics
    stranger = Client(server="127.0.0.1", authport=18121, acctport=18131, secret=b"other",
                      dict=Dictionary(config.RADIUS_DICTIONARY))
    stranger.timeout, stranger.retries = 1, 1
    with pytest.raises(Exception):
        acct(stranger, 1, "X1", "192.168.88.30", ident="Office-Shiraz")
    d = radius_server.diag[("127.0.0.1", "Office-Shiraz", "")]
    assert d["bad_secret"] == 1
    assert radius_server.diag[("127.0.0.1", "Office-Tehran", "")]["name"] == "tehran"
    db = SessionLocal()
    s = db.query(AcctSession).filter_by(acct_session_id="NAT1").one()
    assert s.nas_ip == "10.20.0.1" and s.active
    sara_id = s.user_id
    ali_id = db.query(User).filter_by(username="hq_staff").one().id
    db.add(AcctSession(acct_session_id="HQ1", nas_ip="127.0.0.1", username="hq_staff", user_id=ali_id,
                       framed_ip="192.168.88.20", start_time=s.start_time, last_update=s.last_update, active=True))
    db.add(AcctSession(acct_session_id="HQ2", nas_ip="127.0.0.1", username="hq_staff", user_id=ali_id,
                       framed_ip="192.168.88.77", start_time=s.start_time, last_update=s.last_update, active=True))
    db.commit()
    db.close()

    # Traffic-Flow through the same NAT address
    col = Collector()
    hdr = struct.pack("!HHIIIIBBH", 5, 1, 0, int(time.time()), 0, 0, 0, 0, 0)

    def rec(src, dst, octets):
        return (socket.inet_aton(src) + socket.inet_aton(dst) + b"\0" * 8 + struct.pack("!II", 1, octets) +
                b"\0" * 8 + struct.pack("!HH", 50000, 443) + b"\0" + b"\0" + bytes([6]) + b"\0" * 9)
    col.handle(hdr + rec("192.168.88.77", "100.100.1.1", 700), "127.0.0.1", 40001)     # only on the main router
    assert col.per_router[hq_id]["matched"] == 1
    db = SessionLocal()
    db.query(AcctSession).filter_by(acct_session_id="HQ2").delete()
    db.commit()
    db.close()
    col.ip_user_ts = 0
    # 192.168.88.20 is given out by both offices: the shared port goes to the router of the sender address,
    # the branch's own port to the branch
    col.handle(hdr + rec("192.168.88.20", "100.100.1.1", 1000), "127.0.0.1", 40001)
    col.handle(hdr + rec("192.168.88.20", "100.100.1.1", 3000), "127.0.0.1", 40002, rid=bid, label="nat → 20561")
    col.flush()
    db = SessionLocal()
    cats = {(c.user_id, c.category): c.bytes_up for c in db.query(TrafficCategory).filter(
        TrafficCategory.user_id.in_([sara_id, ali_id]))}
    assert cats[(sara_id, "lan")] == 3000                     # 100.100.0.0/16 is inside the branch's network
    assert (ali_id, "lan") not in cats                        # ... but internet for the main office
    assert col.per_router[bid]["matched"] == 1 and col.exporters["nat → 20561"]["routers"] == {bid: 1}
    db.close()
    # a branch user whose address only that branch gives out needs no own port
    db = SessionLocal()
    db.query(AcctSession).filter_by(acct_session_id="HQ1").delete()
    db.commit()
    db.close()
    col.ip_user_ts = 0
    col.handle(hdr + rec("192.168.88.20", "100.100.1.1", 500), "127.0.0.1", 40001)
    assert col.per_router[bid]["matched"] == 2

    # dedicated port opened on save, packets on it belong to that router
    monkeypatch.setattr(config, "NETFLOW_BIND", "127.0.0.1")
    col.started = True
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/nas")
        assert "Office-Shiraz" in r.text and "افزودن این روتر" in r.text and "Office-Tehran" in r.text
        assert 'value="Office-Shiraz"' in c.get("/nas?ident=Office-Shiraz&nas_ip=").text
        r = c.post("/nas/save", data={"id": str(bid), "name": "tehran", "ip": "10.20.0.1", "identifier": "Office-Tehran",
                                      "flow_port": "18121", "lan_ranges": "100.100.0.0/16"}, follow_redirects=True)
        assert "قابل استفاده نیست" in r.text
        c.post("/nas/save", data={"id": str(bid), "name": "tehran", "ip": "10.20.0.1", "identifier": "Office-Tehran",
                                  "flow_port": "20561", "lan_ranges": "100.100.0.0/16"})
    col.sync_ports()
    assert col.ports == {20561: bid}
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    s.sendto(hdr + rec("192.168.88.20", "8.8.8.8", 10), ("127.0.0.1", 20561))
    s.close()
    for _ in range(50):
        if col.per_router[bid]["matched"] == 3:
            break
        time.sleep(0.05)
    assert col.per_router[bid]["matched"] == 3
    acct(bcl, 2, "NAT1", "192.168.88.20", out=10)
    db = SessionLocal()
    db.delete(db.get(Nas, bid))
    db.commit()
    db.close()
    radius_server.nas_cache.invalidate()


def test_active_directory_users_as_panel_users(env, monkeypatch):
    """An office manager signs in to the panel with their AD account, scoped to their own office in AD; the same
    login page sends internet users without panel access to their usage page."""
    from ldap3 import MODIFY_REPLACE

    from app import ldap_ad, web
    from app.db import Office, PanelUser, set_setting
    from app.web import app
    person = {"objectClass": ["top", "person", "user"], "objectCategory": "person", "userAccountControl": 512}
    sara_dn = "CN=Sara T,OU=Staff,DC=corp,DC=local"
    entries = [(sara_dn, {**person, "sAMAccountName": "sara.t", "displayName": "Sara T",
                          "physicalDeliveryOfficeName": "Tehran-Office"}),
               ("CN=Omid S,OU=Staff,DC=corp,DC=local", {**person, "sAMAccountName": "omid.s", "displayName": "Omid S",
                                                        "physicalDeliveryOfficeName": "Tehran-Office"}),
               ("CN=Kian P,OU=Staff,DC=corp,DC=local", {**person, "sAMAccountName": "kian.p", "displayName": "Kian P",
                                                        "physicalDeliveryOfficeName": "Shiraz-Office"})]
    conn = _mock_ad(entries)
    conn.unbind = lambda *a, **k: True
    monkeypatch.setattr(ldap_ad, "_connect", lambda c, timeout=30: conn)
    passwords = {"sara.t": "AdPass1", "omid.s": "AdPass2"}
    monkeypatch.setattr(ldap_ad, "authenticate", lambda db, u, p: (True, "") if passwords.get(u.lower()) == p
                        else (False, "AD: invalidCredentials"))
    db = SessionLocal()
    for k, v in {"ldap_server": "dc1", "ldap_domain": "corp.local", "ldap_bind_user": "svc",
                 "ldap_base_dn": "DC=corp,DC=local", "ldap_user_filter": "(sAMAccountName=*.*)",
                 "ldap_default_group": ""}.items():
        set_setting(db, k, v)
    db.commit()
    ldap_ad.sync_users(db)
    sara = db.query(User).filter_by(username="sara.t").one()
    sara_id, tehran = sara.id, sara.office_id
    shiraz = db.query(Office).filter_by(name="Shiraz-Office").one().id
    admin_id = db.query(PanelUser).filter_by(username="admin").one().id
    db.close()
    web.MY_FAILS.clear()
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        assert f"/admins/new?user={sara_id}" in c.get(f"/users/{sara_id}/edit").text
        r = c.get(f"/admins/new?user={sara_id}")
        assert 'value="sara.t"' in r.text and 'id="own_office" checked' in r.text
        r = c.post("/admins/save", data={"source": "ldap", "username": "CORP\\sara.t", "role": "unit_manager",
                                         "scope_own_office": "1", "enabled": "1", "full_name": ""},
                   follow_redirects=True)
        assert "کاربر پنل ذخیره شد" in r.text and "AD" in r.text
        # the installer's local admin cannot be turned into an AD account: the panel must work without AD
        r = c.post("/admins/save", data={"id": str(admin_id), "source": "ldap", "username": "admin", "role": "admin",
                                         "enabled": "1"}, follow_redirects=True)
        assert "حساب محلی (غیر AD)" in r.text
    db = SessionLocal()
    p = db.query(PanelUser).filter_by(username="sara.t").one()
    assert p.source == "ldap" and p.user_id == sara_id and p.scope_offices == str(tehran) and p.full_name == "Sara T"
    assert db.get(PanelUser, admin_id).source == "local"
    db.close()
    with TestClient(app) as c:
        r = c.post("/login", data={"username": "sara.t", "password": "wrong"}, follow_redirects=True)
        assert "اشتباه" in r.text
        r = c.post("/login", data={"username": "sara.t@corp.local", "password": "AdPass1"}, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] != "/account"
        r = c.get("/users")
        assert "omid.s" in r.text and "kian.p" not in r.text
        assert "Active Directory" in c.get("/account").text
        assert "رمز Active Directory" in c.post("/account", data={"current": "x", "new": "Aa123456", "again": "Aa123456"},
                                                follow_redirects=True).text
        assert "سلام" in c.get("/my").text                    # same person: their own usage without a second login
        assert c.get("/admins").status_code == 403
    # office changed in AD -> scope follows on the next sync
    conn.modify(sara_dn, {"physicalDeliveryOfficeName": [(MODIFY_REPLACE, ["Shiraz-Office"])]})
    db = SessionLocal()
    ldap_ad.sync_users(db)
    assert db.query(PanelUser).filter_by(username="sara.t").one().scope_offices == str(shiraz)
    db.close()
    with TestClient(app) as c:
        # an internet user without panel access: the same login page opens their usage page
        r = c.post("/login", data={"username": "omid.s", "password": "AdPass2"}, follow_redirects=False)
        assert r.headers["location"] == "/my"
        assert "سلام" in c.get("/my").text
        assert c.get("/users", follow_redirects=False).headers["location"] == "/login"
    db = SessionLocal()
    db.query(PanelUser).filter_by(username="sara.t").delete()
    for k in ("ldap_server", "ldap_base_dn", "ldap_user_filter"):
        set_setting(db, k, "")
    db.commit()
    db.close()
    web.MY_FAILS.clear()


def test_router_reuses_active_id_for_a_new_login(env, monkeypatch):
    """A hotspot login's .id comes from its IP: the same user logging in again from the same address (or the
    counters starting over) gets the id of the earlier, finished session. Seen on a real router as
    'UNIQUE constraint failed: sessions.nas_ip, sessions.acct_session_id, sessions.username'."""
    from app import jobs, routeros
    from tests.fake_ros import FakeRouter
    fake = FakeRouter()
    entry = {".id": "*C0A86F16", "user": "user2", "address": "192.168.111.22", "mac-address": "E8:9C:25:28:57:B7",
             "uptime": "5m", "bytes-in": "1000", "bytes-out": "9000"}
    fake.store["ip/hotspot/active"] = [dict(entry)]
    monkeypatch.setattr(routeros, "api", fake.conn)
    db = SessionLocal()
    nas = Nas(name="reuse", ip="192.168.111.1", secret="s", api_user="api")
    db.add(nas)
    db.commit()
    nid = nas.id
    db.close()
    sid = f"api:{nid}:hs:*C0A86F16"

    def poll():
        jobs.live["routers"].pop("reuse", None)
        jobs.poll_routers()
        assert jobs.live["routers"]["reuse"].startswith("OK"), jobs.live["routers"]["reuse"]
    poll()
    fake.store["ip/hotspot/active"] = []                                  # logout
    poll()
    fake.store["ip/hotspot/active"] = [dict(entry, uptime="10s", **{"bytes-out": "500"})]   # same login again
    poll()
    fake.store["ip/hotspot/active"] = [dict(entry, uptime="20s", **{"bytes-out": "100"})]   # counters reset
    poll()
    db = SessionLocal()
    rows = db.query(AcctSession).filter(AcctSession.nas_ip == "192.168.111.1", AcctSession.username == "user2") \
        .order_by(AcctSession.id).all()
    assert [(r.active, r.terminate_cause) for r in rows] == [(False, "Logout"), (False, "Replaced"), (True, "")]
    assert rows[-1].acct_session_id == sid and rows[0].acct_session_id.startswith(sid + "~")
    for r in rows:
        db.delete(r)
    db.delete(db.get(Nas, nid))
    db.commit()
    db.close()


def test_first_login_tour(env):
    """Shown automatically once per panel user and per internet user (/my), replayable, reset by an admin."""
    import json
    import re

    from app.db import PanelUser
    from app.web import app

    def steps(html):
        return json.loads(re.search(r'id="tour-steps">(.*?)</script>', html, re.S).group(1))
    db = SessionLocal()
    for p in db.query(PanelUser):
        p.tour_seen = ""
    db.commit()
    db.close()
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/")
        assert "auto: true" in r.text and 'data-tour="nas"' in r.text and "راهنمای برنامه" in r.text
        titles = [s["title"] for s in steps(r.text)]
        assert titles[0].startswith("خوش آمدید") and "روترها" in titles and "کاربران پنل" in titles
        assert c.post("/tour/done").status_code == 204
        r = c.get("/users")
        assert "auto: false" in r.text and 'id="tour-steps"' in r.text      # still replayable from the menu
        assert c.get("/static/tour.js").status_code == 200
        # internet user's own page
        c.post("/my/login", data={"username": "ali", "password": "Pass123"})
        r = c.get("/my")
        assert "auto: true" in r.text and "حجم باقی‌مانده" in [s["title"] for s in steps(r.text)]
        assert c.post("/my/tour/done").status_code == 204
        assert "auto: false" in c.get("/my").text
        # an admin can show it again to everyone
        c.post("/admins/reset-tour")
        assert "auto: true" in c.get("/").text and "auto: true" in c.get("/my").text
    with TestClient(app) as c:                                  # not logged in: nothing to mark
        assert c.post("/tour/done", follow_redirects=False).status_code == 303
        assert c.post("/my/tour/done").status_code == 204


def test_router_traffic_graphs(env, monkeypatch):
    """Interface counters read every poll become per-minute speed and volume; today is drawn over yesterday,
    internet (WAN) and internal network (LAN) apart."""
    from datetime import datetime, timedelta

    from app import ifstats, jobs, routeros
    from app.db import IfaceDaily, IfaceSample, now
    from app.web import app
    from tests.fake_ros import FakeRouter
    fake = FakeRouter()
    fake.store["interface"] = [
        {".id": "*1", "name": "ether1", "type": "ether", "rx-byte": "1000", "tx-byte": "500"},
        {".id": "*2", "name": "bridge-lan", "type": "bridge", "rx-byte": "300", "tx-byte": "900"},
        {".id": "*3", "name": "<l2tp-ali>", "type": "l2tp-in", "dynamic": "true", "rx-byte": "1", "tx-byte": "1"}]
    monkeypatch.setattr(routeros, "api", fake.conn)
    db = SessionLocal()
    nas = Nas(name="graphs", ip="10.77.0.1", secret="s", api_user="api", wan_interfaces="ether1",
              lan_interfaces="bridge-lan")
    db.add(nas)
    db.commit()
    nid = nas.id
    db.close()
    jobs.live["routers"].pop("graphs", None)
    jobs.poll_routers()                                     # first poll only seeds the counters
    assert jobs.live["routers"]["graphs"].startswith("OK") and jobs.live["api_bytes"][nid] > 2000
    db = SessionLocal()
    assert db.query(IfaceSample).filter(IfaceSample.nas_id == nid).count() == 0
    t = now().replace(second=0, microsecond=0)
    if t.hour == 0 and t.minute < 3:
        t += timedelta(minutes=3)
    y = t - timedelta(days=1)
    ifs = lambda wrx, wtx, lrx, ltx: [{"name": "ether1", "rx-byte": wrx, "tx-byte": wtx},     # noqa: E731
                                      {"name": "bridge-lan", "rx-byte": lrx, "tx-byte": ltx},
                                      {"name": "<l2tp-ali>", "type": "l2tp-in", "dynamic": True, "rx-byte": 5, "tx-byte": 5}]
    ifstats._last.clear()
    ifstats.record(db, nid, ifs(0, 0, 0, 0), y - timedelta(minutes=1))
    ifstats.record(db, nid, ifs(6_000_000, 600_000, 500_000, 5_000_000), y)       # yesterday: 6 MB in a minute
    ifstats.record(db, nid, ifs(6_000_000, 600_000, 500_000, 5_000_000), t - timedelta(minutes=2))
    ifstats.record(db, nid, ifs(13_500_000, 1_350_000, 1_000_000, 12_500_000), t - timedelta(minutes=1))
    ifstats.record(db, nid, ifs(10, 10, 10, 10), t)                              # router rebooted: skipped
    db.commit()
    assert {r.iface for r in db.query(IfaceSample).filter(IfaceSample.nas_id == nid)} == {"ether1", "bridge-lan"}
    assert db.query(IfaceDaily).filter(IfaceDaily.nas_id == nid, IfaceDaily.iface == "ether1",
                                       IfaceDaily.day == t.date()).one().rx == 7_500_000
    db.close()
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get(f"/traffic?nas={nid}")
        assert r.status_code == 200 and "اینترنت (WAN)" in r.text and "شبکه داخلی (LAN)" in r.text
        assert 'data-tour="traffic"' in r.text and "bridge-lan" in r.text
        w = c.get(f"/api/traffic?nas={nid}&role=wan").json()
        m = (t - timedelta(minutes=1)).hour * 60 + (t - timedelta(minutes=1)).minute
        assert w["is_today"] and w["cur"]["down"][m] == 1_000_000 and w["cur"]["up"][m] == 100_000   # bit/s
        assert w["cur"]["total_down"] == 7_500_000 and w["prev"]["total_down"] == 6_000_000
        assert w["prev"]["down"][y.hour * 60 + y.minute] == 800_000
        lan = c.get(f"/api/traffic?nas={nid}&role=lan").json()       # towards the users = router's tx
        assert lan["cur"]["total_down"] == 7_500_000 and lan["cur"]["total_up"] == 500_000
        one = c.get(f"/api/traffic?nas={nid}&role=lan&iface=ether1").json()
        assert one["ifaces"] == ["ether1"]
        assert c.get("/api/traffic?nas=99999").status_code == 404
        assert "بار مانیتورینگ روی لینک" in c.get("/nas").text
        # internal-network interfaces are chosen on the router form
        r = c.post("/nas/save", data={"id": nid, "name": "graphs", "ip": "10.77.0.1", "api_user": "api", "wan": "ether1",
                                        "lan_if": "bridge-lan", "lan_if_extra": "vlan10"}, follow_redirects=False)
        assert r.status_code == 303
    db = SessionLocal()
    assert db.get(Nas, nid).lan_interfaces == "bridge-lan,vlan10"
    db.query(IfaceSample).filter(IfaceSample.nas_id == nid).delete()
    db.query(IfaceDaily).filter(IfaceDaily.nas_id == nid).delete()
    db.delete(db.get(Nas, nid))
    db.commit()
    db.close()


def test_attendance_newest_first_today_by_default(env):
    from datetime import timedelta

    from app import mgmt
    from app.db import now
    from app.web import app
    db = SessionLocal()
    rows = mgmt.attendance(db, now() - timedelta(days=30), now())
    db.close()
    assert [(r["day"], r["first"]) for r in rows] == sorted([(r["day"], r["first"]) for r in rows], reverse=True)
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/attendance")
        assert '<option value="today" selected>' in r.text.replace("  ", " ")


def test_ldaps_to_old_domain_controller(tmp_path):
    """Windows Servers with only TLS 1.0 and a self-signed certificate: Python's default context fails
    (Schannel drops it: WinError 10054); the app's own TLS settings connect."""
    import datetime
    import socket
    import ssl
    import threading
    import types
    import warnings

    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.x509.oid import NameOID

    from app.ldap_ad import _DCTls
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "Srv-DC.BC.local")])
    t = datetime.datetime.now(datetime.timezone.utc)
    cert = x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key()) \
        .serial_number(1).not_valid_before(t).not_valid_after(t + datetime.timedelta(days=1)).sign(key, hashes.SHA256())
    (tmp_path / "c.pem").write_bytes(cert.public_bytes(serialization.Encoding.PEM))
    (tmp_path / "k.pem").write_bytes(key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                                       serialization.NoEncryption()))
    srv = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    srv.load_cert_chain(tmp_path / "c.pem", tmp_path / "k.pem")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", DeprecationWarning)
            srv.minimum_version = srv.maximum_version = ssl.TLSVersion.TLSv1
        srv.set_ciphers("DEFAULT:@SECLEVEL=0")
    except (ValueError, ssl.SSLError):
        pytest.skip("this OpenSSL has no TLS 1.0")
    ls = socket.socket()
    ls.bind(("127.0.0.1", 0))
    ls.listen(2)

    def serve():
        for _ in range(2):
            c, _ = ls.accept()
            try:
                srv.wrap_socket(c, server_side=True).recv(1)
            except Exception:
                pass
    threading.Thread(target=serve, daemon=True).start()
    port = ls.getsockname()[1]
    default = ssl.create_default_context()
    default.check_hostname, default.verify_mode = False, ssl.CERT_NONE
    with pytest.raises(ssl.SSLError):
        default.wrap_socket(socket.create_connection(("127.0.0.1", port)))
    conn = types.SimpleNamespace(socket=socket.create_connection(("127.0.0.1", port)))
    _DCTls().wrap_socket(conn, do_handshake=True)
    assert conn.socket.version() == "TLSv1"
    ls.close()


def test_issabel_calls_two_servers(env, monkeypatch):
    """CDR rows of two Issabel servers: direction, extension owners by period (extension changes hands, a person
    with two extensions), calls between the two servers, missed calls nobody returned, reports and Excel."""
    from datetime import datetime, timedelta
    from io import BytesIO

    from openpyxl import load_workbook

    from app import jalali, pbx
    from app.db import Call, ExtensionAssignment, PbxServer, now
    from app.web import app
    db = SessionLocal()
    for m in (Call, ExtensionAssignment, PbxServer):
        db.query(m).delete()
    a = PbxServer(name="ایزابل مرکزی", host="10.0.0.5", db_user="r", ext_ranges="100-399", enabled=True)
    b = PbxServer(name="ایزابل تهران", host="10.0.1.5", db_user="r", ext_ranges="500-599", enabled=True)
    db.add_all([a, b])
    ali = db.query(User).filter(User.username == "ali").one()
    sara = User(username="sara.calls", full_name="سارا")
    reza = User(username="reza.calls", full_name="رضا")
    db.add_all([sara, reza])
    db.commit()
    aid, bid, ali_id, sara_id, reza_id = a.id, b.id, ali.id, sara.id, reza.id
    t = datetime.combine(now().date(), datetime.min.time()) + timedelta(hours=9)
    if t > now():
        t -= timedelta(days=1)
    day = t.date()
    db.close()

    def row(m, uid, src, dst, ch, dch, disp="ANSWERED", dur=60, bill=50, seq=0, linked=None, d=t):
        return {"calldate": d + timedelta(minutes=m), "src": src, "dst": dst, "channel": ch, "dstchannel": dch,
                "disposition": disp, "duration": dur, "billsec": bill if disp == "ANSWERED" else 0,
                "uniqueid": uid, "linkedid": linked or uid, "sequence": seq}
    yesterday = t - timedelta(days=1)
    rows = {aid: [
        row(0, "a1", "101", "09121112233", "SIP/101-0001", "SIP/tci-0002"),                          # out, ali
        row(5, "a2", "02155554444", "101", "SIP/tci-0003", "SIP/101-0004", dur=40, bill=30),         # in answered
        row(7, "a3", "02166667777", "101", "SIP/tci-0005", "Local/101@from-queue-0006;1", "NO ANSWER", 25, 0, 1),
        row(7, "a3", "02166667777", "102", "SIP/tci-0005", "Local/102@from-queue-0007;1", "NO ANSWER", 25, 0, 2),
        row(9, "a4", "02188889999", "s", "SIP/tci-0008", "", "NO ANSWER", 10, 0),                    # missed, called back
        row(20, "a5", "101", "02188889999", "SIP/101-0009", "SIP/tci-0010", bill=100),
        row(30, "a6", "101", "102", "SIP/101-0011", "SIP/102-0012", bill=20),                        # internal
        row(40, "a7", "102", "510", "SIP/102-0013", "SIP/link-tehran-0014", bill=15),                # to other server
        row(50, "a8", "250", "00493012345", "SIP/250-0015", "SIP/tci-0016", bill=200),               # reza (2nd ext)
        row(0, "a0", "102", "09350000000", "SIP/102-0017", "SIP/tci-0018", bill=70, d=yesterday),   # 102 = sara then
    ], bid: [
        row(41, "b1", "102", "510", "SIP/link-hq-0001", "SIP/510-0002", bill=15),                   # same call, srv B
        row(45, "b2", "02133332222", "510", "SIP/mokhaberat-01", "SIP/510-0003", bill=90),
    ]}
    calls = {"n": 0}

    def fake_fetch(server, since, limit=pbx.BATCH):
        calls["n"] += 1
        return sorted([r for r in rows[server.id] if r["calldate"] >= since], key=lambda r: r["calldate"])[:limit]
    monkeypatch.setattr(pbx, "fetch", fake_fetch)
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        # extension owners: 101 ali; 102 sara until yesterday, reza from today; reza also has 250
        for uid, ext, start, end in ((ali_id, "101", "", ""), (sara_id, "102", "", ""), (reza_id, "250", "", "")):
            c.post(f"/users/{uid}/ext", data={"extension": ext, "start": start, "end": end})
        r = c.post(f"/users/{reza_id}/ext", data={"extension": "102", "start": jalali.to_jalali_str(day)},
                   follow_redirects=False)
        assert r.status_code == 303
        db = SessionLocal()
        sara_102 = db.query(ExtensionAssignment).filter_by(user_id=sara_id, extension="102").one()
        assert sara_102.end == day - timedelta(days=1)                       # handed over: sara's period closed
        db.close()
        assert "تخصیص داده شده است" in c.post(f"/users/{ali_id}/ext", data={"extension": "102"}).text   # overlap
        pbx.sync_all()
        pbx.sync_all()                                                        # re-reading the overlap adds nothing
        db = SessionLocal()
        assert db.query(Call).count() == 12
        by = {x.key.split(":")[0] + ":" + x.key.split(":")[1]: x for x in db.query(Call)}
        assert (by["a1:0"].direction, by["a1:0"].extension, by["a1:0"].user_id) == ("out", "101", ali_id)
        assert (by["a2:0"].direction, by["a2:0"].other) == ("in", "02155554444")
        assert by["a7:0"].direction == "internal" and by["b1:0"].direction == "internal"
        assert by["a0:0"].user_id == sara_id and by["a6:0"].direction == "internal"
        assert by["a3:2"].user_id == reza_id and by["a8:0"].user_id == reza_id
        assert db.get(PbxServer, aid).cursor is not None and not db.get(PbxServer, aid).last_error
        start = datetime.combine(day, datetime.min.time())
        rep = pbx.build(db, start, start + timedelta(days=1))
        assert rep["stats"]["in"]["count"] == 4 and rep["stats"]["in"]["answered"] == 2       # a2, a3, a4, b2
        assert rep["stats"]["out"]["count"] == 3 and rep["stats"]["internal"]["count"] == 3
        assert [m["number"] for m in rep["missed"]] == ["02166667777"]       # a4 was called back by a5
        assert rep["destinations"][0]["category"] in ("موبایل", "شهری", "بین‌الملل")
        people = {p["user_id"]: p for p in rep["people"]}
        assert people[ali_id]["out"] == 2 and people[ali_id]["in_answered"] == 1 and people[ali_id]["missed"] == 1
        assert people[reza_id]["exts"] == "102، 250" and people[reza_id]["missed"] == 1
        assert sara_id not in people
        assert {u["extension"] for u in rep["unassigned"]} == {"510"}
        db.close()
        r = c.get("/calls?preset=today" if day == now().date() else "/calls?preset=yesterday")
        assert r.status_code == 200 and "تماس‌های بی‌پاسخی که کسی برنگشته" in r.text and "02166667777" in r.text
        assert "داخلی‌های بدون صاحب" in r.text and 'data-tour="calls"' in r.text
        assert c.get("/calls?dir=in").status_code == 200
        wb = load_workbook(BytesIO(c.get("/calls/export.xlsx?preset=30d").content))
        assert "تماس افراد" in wb.sheetnames and "ریز تماس‌ها" in wb.sheetnames
        assert "تماس‌های تلفنی (ایزابل)" in c.get("/mgmt?preset=30d").text
        assert "تماس افراد" in load_workbook(BytesIO(c.get("/mgmt/export.xlsx?preset=30d").content)).sheetnames
        r = c.get("/pbx")
        assert "ایزابل تهران" in r.text and "510" in r.text
        # quick assignment from the unassigned list relinks past calls
        c.post("/pbx/assign", data={"extension": "510", "pbx_id": bid, "user_id": sara_id})
        db = SessionLocal()
        assert {x.user_id for x in db.query(Call).filter(Call.pbx_id == bid)} == {sara_id}
        db.close()
        assert c.get(f"/users/{reza_id}/edit").text.count("/ext/") >= 2
    db = SessionLocal()
    for m in (Call, ExtensionAssignment, PbxServer):
        db.query(m).delete()
    db.query(User).filter(User.id.in_([sara_id, reza_id])).delete()
    db.commit()
    db.close()


def test_pbx_parsing():
    from types import SimpleNamespace

    from app import pbx
    assert pbx.peer("SIP/tci-trunk-0000b2") == "tci-trunk" and pbx.peer("Local/101@from-queue-0001;2") == "101"
    assert pbx.peer("PJSIP/200-00000012") == "200" and pbx.parse_ranges("100-399, 5000") == [(100, 399), (5000, 5000)]
    assert pbx.category("+989121234567") == "موبایل" and pbx.category("0313456789") == "بین‌شهری"
    assert pbx.category("0049301234") == "بین‌الملل" and pbx.category("88776655") == "شهری"
    cl = pbx.Classifier(SimpleNamespace(ext_ranges="100-399", trunks="tci"), [[(100, 399)]])
    assert cl({"channel": "SIP/101-1", "dstchannel": "SIP/tci-2", "src": "101", "dst": "0912"}) == ("out", "101", "0912")
    assert cl({"channel": "SIP/tci-1", "dstchannel": "", "src": "0218", "dst": "s"}) == ("in", "", "0218")


def test_equipment_types_logons_and_nightly_list(env, monkeypatch):
    """Printers / phones / thin clients recognised and set by hand (one or many), Windows logons from the DC
    Security log (IP -> MAC from device history) and from the logon script, and the 21:00 list of devices that are on."""
    from datetime import datetime, timedelta
    from io import BytesIO

    from openpyxl import load_workbook

    from app import devices, logons, nightlist
    from app.db import Device, NightList, get_setting, now, set_setting
    from app.web import app

    assert devices.classify("AA:BB:CC:00:00:91", hostname="NPI3A2B1C")["device_type"] == "printer"
    assert devices.classify("AA:BB:CC:00:00:92", hostname="SEP001122334455", dhcp_class="MSFT 5.0")["device_type"] == "phone"
    assert devices.classify("AA:BB:CC:00:00:93", hostname="WYSE-ACC-12", dhcp_class="MSFT 5.0")["device_type"] == "thin"
    assert devices.classify("AA:BB:CC:00:00:94", hostname="DESKTOP-AB12", dhcp_class="MSFT 5.0")["device_type"] == "pc"

    db = SessionLocal()
    ts = now()
    t0 = ts - timedelta(minutes=30)
    pc = devices.record(db, "AA:BB:DD:00:00:01", ip="192.168.88.21", nas_id=None, source="arp", ts=t0,
                        hostname="ACC-PC1")
    devices.record(db, "AA:BB:DD:00:00:01", ip="192.168.88.21", source="arp", ts=ts)
    devices.record(db, "AA:BB:DD:00:00:01", ip="192.168.88.21", user="ali.hs", source="hotspot", ts=ts)
    devices.record(db, "AA:BB:DD:00:00:02", user="old.hs", source="hotspot", ts=ts - timedelta(hours=2))
    devices.record(db, "AA:BB:DD:00:00:02", ip="192.168.88.22", source="arp", ts=ts, hostname="THIN-02")
    devices.record(db, "AA:BB:DD:00:00:03", ip="192.168.88.30", source="arp", ts=ts, hostname="NPI0A0B0C")
    devices.record(db, "AA:BB:DD:00:00:04", ip="192.168.88.40", source="arp", ts=ts - timedelta(hours=3))  # off
    db.commit()
    ids = {d.mac: d.id for d in db.query(Device).filter(Device.mac.like("AA:BB:DD:%"))}
    assert pc.device_type == "unknown" or pc.device_type

    # DC Security log: TGT for "BC\\ali" from the PC's IP, a machine account and a failed request are ignored
    out = "\n".join([
        f"{(t0 + timedelta(minutes=5)):%Y-%m-%dT%H:%M:%S}\t4768\tali\t::ffff:192.168.88.21\t0x0",
        f"{(t0 + timedelta(minutes=6)):%Y-%m-%dT%H:%M:%S}\t4768\tACC-PC1$\t::ffff:192.168.88.21\t0x0",
        f"{(t0 + timedelta(minutes=7)):%Y-%m-%dT%H:%M:%S}\t4768\treza\t::ffff:192.168.88.21\t0x18",
        f"{(t0 + timedelta(minutes=8)):%Y-%m-%dT%H:%M:%S}\t4768\tnobody\t10.9.9.9\t0x0",
    ])
    seen = {}

    def fake_read(dc, account, password, since):
        seen.update(dc=dc, account=account)
        return out
    monkeypatch.setattr(logons, "read_dc", fake_read)
    set_setting(db, "ldap_server", "srv-dc.bc.local")
    set_setting(db, "ldap_domain", "bc.local")
    set_setting(db, "ldap_bind_user", "svc")
    set_setting(db, "logon_dc_enabled", "1")
    db.commit()
    db.close()
    logons.sync_dcs()
    db = SessionLocal()
    assert seen == {"dc": "srv-dc.bc.local", "account": "bc.local\\svc"}
    d1 = db.query(Device).filter(Device.mac == "AA:BB:DD:00:00:01").one()
    assert d1.logon_user == "ali" and d1.logon_source == "dc"
    assert d1.last_user == "ali.hs"                                  # hotspot user kept apart from the Windows logon
    assert get_setting(db, "logon_status").endswith("|2|1")       # 2 successful user logons, 1 matched a device
    set_setting(db, "ldap_server", "")
    db.commit()
    db.close()

    with TestClient(app) as c:
        # logon script: wrong key refused; right key records the user on the computer that sent it (by name here)
        assert c.get("/logon?u=sara&c=THIN-02&k=x").status_code == 403
        _login(c, "admin", "Secret12345")
        c.post("/devices/logon-settings", data={"night_list_time": "۲۱:۰۰", "logon_dc_enabled": "1", "new_key": "1"})
        db = SessionLocal()
        key = get_setting(db, "logon_script_key")
        db.close()
        assert len(key) == 24
        r = c.get(f"/logon?u=BC\\sara&c=THIN-02&k={key}")
        assert r.status_code == 200 and r.text == "ok"
        assert key in c.get("/devices").text
        # equipment types: one by hand with a name, two at once
        c.post(f"/devices/{ids['AA:BB:DD:00:00:03']}/type", data={"device_type": "printer", "label": "پرینتر طبقه ۲"})
        c.post("/devices/bulk-type", data={"device_type": "thin", "ids": [str(ids["AA:BB:DD:00:00:02"])]})
        c.post("/devices/bulk-type", data={"device_type": "pc", "ids": [str(ids["AA:BB:DD:00:00:01"]),
                                                                         str(ids["AA:BB:DD:00:00:04"])]})
        page = c.get("/devices?type=computers").text
        assert "AA:BB:DD:00:00:02" in page and "AA:BB:DD:00:00:03" not in page
        assert "پرینتر طبقه ۲" in c.get("/devices?type=printer").text

        # nightly list: only at/after the configured time, once per night; devices that are on, with the user
        db = SessionLocal()
        db.query(NightList).delete()
        set_setting(db, "night_list_day", "")
        db.commit()
        assert nightlist.check_time(db) == (21, 0)
        db.close()
        r = c.post("/devices/night/take", follow_redirects=False)
        assert r.status_code == 303
        db = SessionLocal()
        rows = {r.mac: r for r in db.query(NightList).all()}
        assert set(rows) >= {"AA:BB:DD:00:00:01", "AA:BB:DD:00:00:02", "AA:BB:DD:00:00:03"}
        assert "AA:BB:DD:00:00:04" not in rows                      # not seen for 3 hours: off
        assert rows["AA:BB:DD:00:00:01"].user == "ali" and rows["AA:BB:DD:00:00:02"].user == "sara"
        assert rows["AA:BB:DD:00:00:02"].user_source == "script" and rows["AA:BB:DD:00:00:03"].label == "پرینتر طبقه ۲"
        assert rows["AA:BB:DD:00:00:01"].hotspot_user == "ali.hs" and rows["AA:BB:DD:00:00:01"].hotspot_active
        assert rows["AA:BB:DD:00:00:02"].hotspot_user == "old.hs" and not rows["AA:BB:DD:00:00:02"].hotspot_active
        db.close()
        page = c.get("/devices/night").text
        assert "ACC-PC1" in page and "THIN-02" in page and "NPI0A0B0C" not in page    # computers by default
        assert "NPI0A0B0C" in c.get("/devices/night?type=").text
        wb = load_workbook(BytesIO(c.get("/devices/night.xlsx?type=").content))
        vals = [str(v) for row in wb.active.iter_rows(values_only=True) for v in row if v]
        assert "ali" in vals and "sara" in vals and "ali.hs" in vals and "old.hs" in vals and "پرینتر" in vals

    # the job: before 21:00 nothing, at 21:05 the list is taken once
    db = SessionLocal()
    db.query(NightList).delete()
    set_setting(db, "night_list_day", "")
    db.commit()
    db.close()
    day = now().date()
    for hhmm, expect in (("20:59", 0), ("21:05", 1), ("21:06", 1)):
        fake = datetime.combine(day, datetime.strptime(hhmm, "%H:%M").time())
        monkeypatch.setattr(nightlist, "now", lambda f=fake: f)
        nightlist.job()
        db = SessionLocal()
        assert (get_setting(db, "night_list_day") == day.isoformat()) == bool(expect)
        db.close()


def test_tunnel_traffic_apart_and_not_from_quota(env):
    """Branch office: traffic to the other offices through a tunnel over the internet is 'tunnel' (tunnel-type
    interfaces found automatically, or the offices' ranges for IPsec), a public address through the tunnel stays
    internet; internal-network and tunnel traffic are not deducted from the hotspot quota."""
    from datetime import datetime, timedelta

    from app import netscope, policy
    from app.db import SpeedSample, TrafficCategory, now, set_setting
    db = SessionLocal()
    br = Nas(name="branch-tun", ip="10.77.0.1", secret="x", wan_interfaces="ether1", tunnel_ranges="172.20.0.0/16")
    db.add(br)
    db.commit()
    nid = br.id
    db.close()
    sc = netscope.scope
    sc.refresh(force=True)
    sc.learn(nid, [{".id": "*1", "name": "ether1", "type": "ether"}, {".id": "*7", "name": "bridge", "type": "bridge"},
                   {".id": "*9", "name": "eoip-hq", "type": "eoip"}])
    assert sc.classify_for(nid, "192.168.1.10", 9, 7) == "tunnel"        # HQ server through the EoIP tunnel
    assert sc.classify_for(nid, "8.8.8.8", 9, 7) == "internet"            # internet that comes from HQ
    assert sc.classify_for(nid, "172.20.4.4", 1, 7) == "tunnel"           # IPsec: leaves through WAN, office range
    assert sc.classify_for(nid, "192.168.50.2", 7, 7) == "lan"            # printer in the same office
    assert sc.classify_for(nid, "1.1.1.1", 1, 7) == "internet"
    assert sc.is_internal_for(nid, "192.168.1.10", 9, 7)
    db = SessionLocal()
    db.get(Nas, nid).tunnel_interfaces = "l2tp-out1"                   # chosen by hand: EoIP no longer a tunnel
    db.commit()
    db.close()
    sc.refresh(force=True)
    assert sc.classify_for(nid, "192.168.1.10", 9, 7) == "lan"

    db = SessionLocal()
    u = User(username="tun.user", group_id=db.query(Group).first().id)
    db.add(u)
    db.commit()
    ts = now()
    MB = 1024 * 1024
    db.add(SpeedSample(user_id=u.id, ts=ts, d_in=40 * MB, d_out=60 * MB))
    db.add(TrafficCategory(user_id=u.id, day=ts.date(), hour=ts.hour, category="lan", bytes_down=20 * MB, bytes_up=10 * MB))
    db.add(TrafficCategory(user_id=u.id, day=ts.date(), hour=ts.hour, category="tunnel", bytes_down=15 * MB, bytes_up=5 * MB))
    db.add(TrafficCategory(user_id=u.id, day=ts.date(), hour=ts.hour, category="video", bytes_down=40 * MB, bytes_up=0))
    db.commit()
    since = datetime.combine(ts.date(), datetime.min.time())
    set_setting(db, "quota_internet_only", "1")           # an earlier test saved the scope form without it
    db.commit()
    assert policy.usage_bytes(db, u, since) == 50 * MB                     # 100 - 30 internal - 20 tunnel
    set_setting(db, "quota_internet_only", "0")
    db.commit()
    assert policy.usage_bytes(db, u, since) == 100 * MB
    set_setting(db, "quota_internet_only", "1")
    db.commit()
    from app.reports import split_usage
    sp = split_usage(db, since, since + timedelta(days=1), [u.id])
    assert (sp["internal"], sp["tunnel"], sp["internet"]) == (30 * MB, 20 * MB, 50 * MB)
    db.query(Nas).filter(Nas.id == nid).delete()
    db.commit()
    db.close()
    sc.refresh(force=True)


def test_locks_grants_logon_message_quota_warning_monitor(env):
    """Failed-login lock (user and MAC) with alert, manual unlock, temporary grants with who ordered them, the Windows
    logon popup text, the 80% quota banner, router / tunnel monitoring and exact internet from WAN counters."""
    from datetime import timedelta

    from app import logons, monitor, policy, security
    from app.db import AcctSession, Alert, IfaceDaily, Lock, NetEvent, TempGrant, now, set_setting
    from app.web import app
    db = SessionLocal()
    set_setting(db, "lock_fail_count", "3")
    g = db.query(Group).first()
    db.add(User(username="lk.user", password_enc=encrypt("Pw123456"), group=g, quota_mb=100))
    db.add(User(username="late.user", password_enc=encrypt("Pw123456"), allowed_hours="fri 03:00-03:01"))
    db.commit()
    db.close()

    def auth(user, pw, mac):
        req = env.CreateAuthPacket(code=packet.AccessRequest, User_Name=user)
        req["User-Password"] = req.PwCrypt(pw)
        req["Calling-Station-Id"] = mac
        return env.SendPacket(req).code

    for _ in range(3):
        assert auth("lk.user", "bad", "AA:BB:EE:00:00:01") == packet.AccessReject
    db = SessionLocal()
    locks = {(lk.kind, lk.value) for lk in security.active_locks(db)}
    assert ("user", "lk.user") in locks and ("mac", "AA:BB:EE:00:00:01") in locks
    assert db.query(Alert).filter(Alert.kind == "lock").count() >= 1
    db.close()
    assert auth("lk.user", "Pw123456", "AA:BB:EE:00:00:02") == packet.AccessReject        # user locked

    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/security")
        assert r.status_code == 200 and "lk.user" in r.text and "AA:BB:EE:00:00:01" in r.text
        db = SessionLocal()
        lid = db.query(Lock).filter(Lock.kind == "user", Lock.value == "lk.user").first().id
        db.close()
        c.post(f"/security/{lid}/unlock")
        assert auth("lk.user", "Pw123456", "AA:BB:EE:00:00:02") == packet.AccessAccept
        assert auth("lk.user", "Pw123456", "AA:BB:EE:00:00:01") == packet.AccessReject   # the MAC is still locked
        c.post("/security/lock", data={"kind": "ip", "value": "10.9.9.9", "reason": "تست"})
        db = SessionLocal()
        assert security.locked(db, "", "10.9.9.9") is not None
        set_setting(db, "lock_fail_count", "10")
        db.commit()
        db.close()
        assert "هشدارها" in c.get("/alerts").text and "قفل خودکار" in c.get("/alerts").text

        # temporary permit: out-of-hours work; "ordered by" is required
        day = __import__("app.jalali", fromlist=["x"]).to_jalali_str(now().date())
        c.post("/grants/add", data={"username": "late.user", "kind": "hours", "start_day": day, "start_time": "00:00",
                                    "end_day": day, "end_time": "23:59", "ordered_by": ""})
        db = SessionLocal()
        assert db.query(TempGrant).count() == 0
        db.close()
        assert auth("late.user", "Pw123456", "AA:BB:EE:00:00:03") == packet.AccessReject
        c.post("/grants/add", data={"username": "late.user", "kind": "hours", "start_day": day, "start_time": "00:00",
                                    "end_day": day, "end_time": "23:59", "ordered_by": "مدیرعامل"})
        assert auth("late.user", "Pw123456", "AA:BB:EE:00:00:03") == packet.AccessAccept
        assert "مدیرعامل" in c.get("/grants").text

        # 80% quota warning: logon popup text, /my banner data and the hotspot banner script
        db = SessionLocal()
        u = db.query(User).filter(User.username == "lk.user").first()
        db.add(SpeedSample(user_id=u.id, ts=now(), d_in=45 * 1024 * 1024, d_out=40 * 1024 * 1024))
        db.add(AcctSession(acct_session_id="w1", nas_ip="127.0.0.1", username="lk.user", user_id=u.id,
                           framed_ip="testclient", active=True, last_update=now()))
        set_setting(db, "logon_script_key", "k123")
        db.commit()
        w = policy.quota_warning(db, u)
        assert w and w["pct"] == 85 and w["level"] == "warn"
        assert logons.message(db, "lk.user") is None                     # popup off by default
        db.close()
        c.post("/devices/logon-settings", data={"night_list_time": "21:00", "logon_msg_form": "1",
                                                "logon_msg_enabled": "1", "logon_message": "سرور فایل امشب در دسترس نیست",
                                                "logon_msg_quota": "warn", "quota_warn_pct": "80", "new_key": ""})
        r = c.get("/logon?m=1&k=k123&u=CORP\\lk.user&c=PC1")
        title, body = r.text.split("\n", 1)
        assert title and "سرور فایل" in body and "۸۵٪" in body
        assert __import__("json").dumps("۸۵٪")[1:-1] in c.get("/hs/warn.js?u=lk.user").text
        assert c.get("/hs/warn.js?u=ali").text == ""                     # no session of that user from this IP

        # monitoring: router up, tunnel down (after 2 polls), ping slow
        db = SessionLocal()
        br = Nas(name="mon-br", ip="10.78.0.1", secret="x", wan_interfaces="ether1", api_user="api")
        db.add(br)
        db.commit()
        snap = {"interfaces": [{"name": "ether1", "type": "ether", "running": "true"},
                               {"name": "eoip-hq", "type": "eoip", "running": "false"}],
                "pings": {"192.168.1.10": [{"seq": "2", "packet-loss": "0", "avg-rtt": "320ms"}]}}
        monitor.observe(db, br, True, snap)
        monitor.observe(db, br, True, snap)
        db.commit()
        states = {e.target: e.state for e in db.query(NetEvent).filter(NetEvent.nas_id == br.id,
                                                                        NetEvent.ended_at.is_(None))}
        assert states == {"router": "up", "tunnel:eoip-hq": "down", "ping:192.168.1.10": "slow"}
        assert br.auto_tunnels == "eoip-hq"
        assert db.query(Alert).filter(Alert.kind == "monitor").count() >= 2
        # exact internet from interface counters: WAN 10 GB, tunnel 3 GB -> 7 GB internet without the tunnel
        GB = 1024 ** 3
        db.add(IfaceDaily(nas_id=br.id, iface="ether1", day=now().date(), rx=8 * GB, tx=2 * GB))
        db.add(IfaceDaily(nas_id=br.id, iface="eoip-hq", day=now().date(), rx=2 * GB, tx=1 * GB))
        db.commit()
        bid = br.id
        db.close()
        r = c.get("/monitor")
        assert r.status_code == 200 and "eoip-hq" in r.text and "قطع" in r.text
        r = c.get("/mgmt?preset=today")
        assert "اینترنت منهای تانل" in r.text and "7.00 GB" in r.text and "10.00 GB" in r.text
    from app import auth as _auth
    from app.db import PanelUser
    db = SessionLocal()
    for name, role in (("mon.it", "it"), ("mon.hr", "hr")):
        db.add(PanelUser(username=name, role=role, enabled=True, password_hash=_auth.hash_password("MonPass2026x")))
    db.commit()
    db.close()
    with TestClient(app) as c:      # IT expert sees monitoring; HR does not
        _login(c, "mon.it", "MonPass2026x")
        assert c.get("/monitor").status_code == 200 and c.get("/security").status_code == 200
    with TestClient(app) as c:
        _login(c, "mon.hr", "MonPass2026x")
        assert c.get("/mgmt").status_code == 200
        for url in ("/monitor", "/security", "/alerts", "/grants"):
            assert c.get(url).status_code == 403, url
    db = SessionLocal()
    db.query(IfaceDaily).filter(IfaceDaily.nas_id == bid).delete()
    db.query(Nas).filter(Nas.id == bid).delete()
    db.query(AcctSession).filter(AcctSession.acct_session_id == "w1").delete()
    db.commit()
    db.close()


def test_issabel_queues_live_calls_ad_suggestions_and_missed_alert(env, monkeypatch):
    """Queue calls from the CDR (Queue() answers the caller, so only an agent leg is 'answered'), wait times and
    service level; live calls from AMI CoreShowChannels; extensions suggested by AD ipPhone; unanswered-call alert."""
    from datetime import timedelta

    from app import ami, pbx
    from app.db import Alert, Call, ExtensionAssignment, PbxServer, now, set_setting
    from app.web import app
    db = SessionLocal()
    for m in (Call, ExtensionAssignment, PbxServer):
        db.query(m).delete()
    s = PbxServer(name="ایزابل صف", host="10.0.0.9", db_user="r", ext_ranges="100-499", enabled=True,
                  ami_user="acct_live", ami_password_enc=encrypt("x"))
    db.add(s)
    db.add(User(username="q.agent", full_name="اپراتور صف", ip_phone="101"))
    db.add(User(username="q.other", full_name="کارمند دیگر", ip_phone="9999"))       # outside the ranges
    set_setting(db, "pbx_missed_alert", "2")
    db.commit()
    sid = s.id
    db.close()
    t = now() - timedelta(minutes=40)

    def row(m, uid, src, dst, ch, dch, disp, dur, bill, app="Dial", ctx="from-internal", seq=0):
        return {"calldate": t + timedelta(seconds=m), "src": src, "dst": dst, "channel": ch, "dstchannel": dch,
                "disposition": disp, "duration": dur, "billsec": bill, "uniqueid": uid, "linkedid": uid,
                "sequence": seq, "lastapp": app, "dcontext": ctx}
    rows = [
        # answered by agent 101 after 12 s of waiting in queue 400
        row(0, "q1", "02155550001", "400", "SIP/tci-01", "SIP/101-02", "ANSWERED", 72, 60, "Queue", "ext-queues"),
        # waited 35 s (outside the 20 s service level), answered by 102
        row(60, "q2", "02155550002", "400", "SIP/tci-03", "SIP/102-04", "ANSWERED", 95, 60, "Queue", "ext-queues"),
        # abandoned after 50 s: Queue() answered the caller, nobody took it
        row(120, "q3", "02155550003", "400", "SIP/tci-05", "", "ANSWERED", 50, 50, "Queue", "ext-queues"),
        row(180, "q4", "02155550004", "400", "SIP/tci-06", "", "ANSWERED", 30, 30, "Queue", "ext-queues"),
    ]
    monkeypatch.setattr(pbx, "fetch", lambda server, since, limit=pbx.BATCH:
                        [r for r in rows if r["calldate"] >= since][:limit])
    pbx.sync_all()
    db = SessionLocal()
    rep = pbx.build(db, t - timedelta(minutes=1), now() + timedelta(minutes=1))
    q = rep["queues"][0]
    assert (q["queue"], q["offered"], q["answered"], q["abandoned"]) == ("400", 4, 2, 2)
    assert (q["avg_wait"], q["max_wait"], q["sl_pct"], q["avg_abandon"]) == (23, 35, 25.0, 40)
    assert rep["stats"]["in"]["count"] == 4 and rep["stats"]["in"]["answered"] == 2      # abandoned = not answered
    al = db.query(Alert).filter(Alert.kind == "calls").one()
    assert "2 تماس ورودی" in al.body and "صف 400: 2 رهاشده" in al.body
    pbx.sync_all()                                                                      # once an hour at most
    assert db.query(Alert).filter(Alert.kind == "calls").count() == 1
    sug = pbx.ad_suggestions(db)
    assert [(g["username"], g["extension"]) for g in sug] == [("q.agent", "101")]
    agent_id = db.query(User).filter(User.username == "q.agent").one().id
    db.close()

    events = [
        {"Event": "CoreShowChannel", "Channel": "SIP/tci-0007", "Uniqueid": "L1", "Linkedid": "L1",
         "CallerIDNum": "09121234567", "ConnectedLineNum": "101", "ChannelStateDesc": "Up", "BridgeId": "b1",
         "Duration": "00:02:05", "Application": "AppQueue"},
        {"Event": "CoreShowChannel", "Channel": "SIP/101-0008", "Uniqueid": "L1.1", "Linkedid": "L1",
         "CallerIDNum": "101", "ConnectedLineNum": "09121234567", "ChannelStateDesc": "Up", "BridgeId": "b1",
         "Duration": "00:01:50", "Application": "AppDial"},
        {"Event": "CoreShowChannel", "Channel": "SIP/tci-0009", "Uniqueid": "L2", "Linkedid": "L2",
         "CallerIDNum": "02144443333", "ChannelStateDesc": "Up", "Duration": "00:00:18", "Application": "Queue",
         "Exten": "400"},
        {"Event": "CoreShowChannel", "Channel": "SIP/102-0010", "Uniqueid": "L3", "Linkedid": "L3",
         "CallerIDNum": "102", "ConnectedLineNum": "02188887777", "ChannelStateDesc": "Ring", "Duration": "00:00:04",
         "Application": "Dial"},
    ]
    monkeypatch.setattr(ami, "channels", lambda server, timeout=6: events)
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        c.post("/pbx/ad-assign", data={})
        db = SessionLocal()
        assert db.query(ExtensionAssignment).filter_by(user_id=agent_id, extension="101").count() == 1
        db.close()
        d = c.get("/api/calls/live").json()
        got = {(r["direction"], r["extension"], r["state"], r["other"], r["queue"]) for r in d["rows"]}
        assert got == {("in", "101", "talk", "09121234567", ""), ("in", "", "wait", "02144443333", "400"),
                       ("out", "102", "ring", "02188887777", "")}
        assert next(r for r in d["rows"] if r["extension"] == "101")["user"] == "اپراتور صف"
        r = c.get("/calls/live")
        assert r.status_code == 200 and "اپراتور صف" in r.text and "انتظار در صف" in r.text
        r = c.get("/calls?preset=today")
        assert "صف‌ها" in r.text and "25.0٪" in r.text
        assert c.get("/pbx").status_code == 200
        c.post("/pbx/settings", data={"pbx_sl_seconds": "40", "pbx_missed_alert": "0"})
        db = SessionLocal()
        assert pbx.build(db, t - timedelta(minutes=1), now() + timedelta(minutes=1))["queues"][0]["sl_pct"] == 50.0
        db.query(PbxServer).filter(PbxServer.id == sid).delete()
        db.query(Call).delete()
        db.commit()
        db.close()


def test_office_report_by_device_and_network_objects(env):
    """Branch without hotspot users: every device's traffic from NetFlow by MAC (with its Windows user), split into
    internet / tunnel / LAN and the defined servers, cameras and phones; a camera sending to HQ over the tunnel."""
    from datetime import timedelta

    from app import netobjects, netscope
    from app.db import Device, DeviceTraffic, IfaceDaily, NetObject, now
    from app.netflow import Collector
    from app.web import app
    db = SessionLocal()
    br = Nas(name="شعبه بدون هات‌اسپات", ip="10.66.0.1", secret="x", wan_interfaces="ether1",
             lan_ranges="192.168.66.0/24", tunnel_ranges="192.168.1.0/24")
    db.add(br)
    db.commit()
    nid = br.id
    db.add_all([
        NetObject(name="سرور فایل مرکزی", kind="server", addresses="192.168.1.10", reach_tunnel=True),
        NetObject(name="ایزابل", kind="pbx", addresses="192.168.1.5"),
        NetObject(name="دوربین‌های شعبه", kind="camera", nas_id=nid, addresses="192.168.66.200-192.168.66.210"),
        NetObject(name="NVR مرکزی", kind="nvr", addresses="192.168.1.20", reach_external=False),
        Device(mac="AA:BB:F0:00:00:01", last_ip="192.168.66.15", last_nas_id=nid, device_type="pc",
               hostname="PC-ACC1", logon_user="m.karimi", last_seen=now()),
        Device(mac="AA:BB:F0:00:00:02", last_ip="192.168.66.201", last_nas_id=nid, device_type="camera",
               last_seen=now()),
        Device(mac="AA:BB:F0:00:00:03", last_ip="192.168.66.40", last_nas_id=nid, device_type="phone",
               last_seen=now()),
        IfaceDaily(nas_id=nid, iface="ether1", day=now().date(), rx=900 * 1024 ** 2, tx=300 * 1024 ** 2),
    ])
    db.commit()
    db.close()
    netscope.scope.refresh(force=True)
    netobjects.index.refresh(force=True)
    MB = 1024 * 1024
    col = Collector()
    ts = now()
    netobjects.index.flush()                                   # flows of earlier tests' routers
    db = SessionLocal()
    db.query(DeviceTraffic).delete()
    db.commit()
    db.close()

    def flow(src, dst, b, dport=443):
        col.add_flow({"src": src, "dst": dst, "bytes": b, "sport": 50000, "dport": dport, "proto": 6, "in_if": 0,
                      "out_if": 0}, ts, rid=nid)
    flow("192.168.66.15", "8.8.8.8", 10 * MB)                  # PC -> internet
    flow("142.250.1.1", "192.168.66.15", 40 * MB)              # internet -> PC
    flow("192.168.66.15", "192.168.1.10", 25 * MB, 445)        # PC -> HQ file server (tunnel)
    flow("192.168.1.10", "192.168.66.15", 5 * MB, 50000)
    flow("192.168.66.15", "192.168.66.7", 3 * MB)              # PC -> printer in the office
    flow("192.168.66.201", "192.168.1.20", 60 * MB, 554)       # camera -> HQ NVR over the tunnel
    flow("192.168.66.201", "1.2.3.4", 2 * MB)                  # camera -> internet (not allowed)
    flow("192.168.66.40", "192.168.1.5", 4 * MB, 5060)         # phone -> Issabel
    flow("10.9.9.9", "8.8.4.4", 99 * MB)                       # not this office: ignored
    netobjects.index.flush()
    db = SessionLocal()
    rep = netobjects.report(db, ts - timedelta(hours=1), ts + timedelta(hours=1), [nid])
    groups = {g["key"]: g["total"] for g in rep["groups"]}
    assert groups == {"internet": 52 * MB, "tunnel": 0, "lan": 3 * MB}
    assert rep["objects_total"] == 94 * MB and rep["total"] == 149 * MB
    pc = next(e for e in rep["devices"] if e["mac"] == "AA:BB:F0:00:00:01")
    assert (pc["internet"], pc["objects"], pc["lan"], pc["windows_user"]) == (50 * MB, 30 * MB, 3 * MB, "m.karimi")
    objs = {x["obj"].name: x for x in rep["objects"]}
    assert objs["سرور فایل مرکزی"]["total"] == 30 * MB and objs["سرور فایل مرکزی"]["devices"] == 1
    assert objs["دوربین‌های شعبه"]["total"] == 62 * MB and objs["دوربین‌های شعبه"]["internet"] == 2 * MB
    assert objs["NVR مرکزی"]["total"] == 60 * MB and objs["ایزابل"]["total"] == 4 * MB
    kinds = {k["kind"]: k["total"] for k in rep["kinds"]}
    assert kinds["camera"] == 62 * MB and kinds["pbx"] == 4 * MB and kinds["phone"] == 4 * MB
    db.close()
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get(f"/offices?nas={nid}&preset=today")
        assert r.status_code == 200 and "m.karimi" in r.text and "دوربین‌های شعبه" in r.text and "1.17 GB" in r.text
        assert c.get("/offices/export.xlsx?preset=today").status_code == 200
        assert c.get("/netobjects").status_code == 200
        r = c.post("/netobjects/save", data={"name": "سرور مالی", "kind": "server", "addresses": "192.168.1.30\nbad",
                                             "nas_id": ""}, follow_redirects=True)
        assert "نادرست" in r.text
        c.post("/netobjects/save", data={"name": "سرور مالی", "kind": "server", "addresses": "192.168.1.30",
                                         "reach_tunnel": "1"})
        assert "سرور مالی" in c.get("/netobjects").text
    db = SessionLocal()
    db.query(DeviceTraffic).delete()
    db.query(NetObject).delete()
    db.query(IfaceDaily).filter(IfaceDaily.nas_id == nid).delete()
    db.query(Device).filter(Device.mac.like("AA:BB:F0:%")).delete(synchronize_session=False)
    db.query(Nas).filter(Nas.id == nid).delete()
    db.commit()
    db.close()
    netscope.scope.refresh(force=True)
    netobjects.index.refresh(force=True)


def test_history_retention_purge_and_routeros_update(env, monkeypatch, tmp_path):
    from datetime import timedelta

    from app import hotspot, retention, rosadmin, routeros, web
    from app.db import AcctSession, AuditLog, SpeedSample, User, now, set_setting
    from app.web import app
    from tests.fake_ros import FakeRouter
    db = SessionLocal()
    g_total = Group(name="ret-total", quota_mb=1000, quota_period="total")
    db.add(g_total)
    db.flush()
    keep_u = User(username="ret.total", source="local", group_id=g_total.id)
    u = User(username="ret.user", source="local")
    db.add_all([keep_u, u])
    db.flush()
    t = now()
    old, mid = t - timedelta(days=250), t - timedelta(days=100)
    for usr in (u, keep_u):
        for ts in (old, mid):
            s = AcctSession(acct_session_id=f"r{usr.id}-{ts.toordinal()}", nas_ip="10.9.9.9", username=usr.username,
                            user_id=usr.id, start_time=ts, stop_time=ts, active=False)
            db.add(s)
            db.flush()
            db.add(SpeedSample(session_id=s.id, user_id=usr.id, ts=ts, seconds=60, d_in=1000, d_out=5000))
    set_setting(db, "history_keep_months", "6")
    db.commit()
    assert retention.keep_months(db) == 6
    retention.auto(db)                                   # older than 6 months goes, except the "total" quota user
    q = lambda uid: db.query(SpeedSample).filter(SpeedSample.user_id == uid).count()   # noqa: E731
    assert q(u.id) == 1 and q(keep_u.id) == 2
    assert db.query(AcctSession).filter(AcctSession.user_id == u.id).count() == 1
    ctx = web._my_context(db, u)
    assert len(ctx["history"]) == 6 and all("total" in m for m in ctx["history"])
    uid, kid = u.id, keep_u.id
    db.close()

    fake = FakeRouter()
    fake.store["system/resource"] = [{"version": "7.15.3 (stable)", "architecture-name": "arm64", "board-name": "RB5009",
                                      "uptime": "3d4h", "cpu-load": "4", "free-memory": "800000000",
                                      "total-memory": "1000000000", "free-hdd-space": "900000000",
                                      "total-hdd-space": "1000000000"}]
    fake.store["system/routerboard"] = [{"routerboard": "true", "model": "RB5009UG+S+", "current-firmware": "7.15.3",
                                         "upgrade-firmware": "7.16"}]
    fake.store["system/package/update"] = [{".id": "*0", "channel": "stable", "installed-version": "7.15.3",
                                            "latest-version": "7.16.2", "status": "New version is available"}]
    fake.store["system/package"] = [{".id": "*1", "name": "routeros", "version": "7.15.3"}]
    monkeypatch.setattr(routeros, "api", fake.conn)
    sent = {}

    class Tr:
        def put(self, path, data):
            sent[path] = data

        def close(self):
            pass
    monkeypatch.setattr(hotspot, "transfer", lambda nas: Tr())
    monkeypatch.setattr(rosadmin, "NPK_DIR", str(tmp_path / "npk"))
    db = SessionLocal()
    nas = Nas(name="ros-upd", ip="10.66.0.1", secret="s", api_user="api")
    db.add(nas)
    db.commit()
    nid = nas.id
    db.close()
    npk = b"\x1e\xf1\xd0\xba" + b"x" * 100
    with TestClient(app) as c:
        _login(c, "admin", "Secret12345")
        r = c.get("/data")
        assert r.status_code == 200 and "نگهداری سابقه" in r.text and "نمونه‌های مصرف" in r.text
        r = c.get("/routeros?check=1")
        assert r.status_code == 200 and "7.15.3" in r.text and "7.16.2" in r.text and "RB5009UG+S+" in r.text
        assert "arm64" in r.text
        # upload: wrong name / not a package / ok
        c.post("/routeros/upload", files={"file": ("evil.exe", b"MZ", "application/octet-stream")})
        c.post("/routeros/upload", files={"file": ("routeros-7.16.2-arm64.npk", b"not a npk", "application/octet-stream")})
        assert rosadmin.list_npk() == []
        c.post("/routeros/upload", files={"file": ("routeros-7.16.2-arm64.npk", npk, "application/octet-stream")})
        c.post("/routeros/upload", files={"file": ("routeros-7.16.2-mipsbe.npk", npk, "application/octet-stream")})
        assert [p["arch"] for p in rosadmin.list_npk()] == ["arm64", "mipsbe"]
        # without confirmation nothing happens
        c.post("/routeros/action", data={"action": "reboot", "nas": nid})
        assert not any(cmd == "/system/reboot" for cmd, _ in fake.calls)
        c.post("/routeros/action", data={"action": "send", "nas": nid, "npk": "routeros-7.16.2-mipsbe.npk", "confirm": "1"})
        assert sent == {}                                    # architecture differs
        c.post("/routeros/action", data={"action": "send", "nas": nid, "npk": "routeros-7.16.2-arm64.npk",
                                         "confirm": "1", "reboot_after": "1"})
        assert sent == {"routeros-7.16.2-arm64.npk": npk}
        assert any(cmd == "/system/reboot" for cmd, _ in fake.calls)
        c.post("/routeros/action", data={"action": "check", "nas": nid, "channel": "long-term"})
        assert ("/system/package/update/set", {"channel": "long-term"}) in fake.calls
        assert any(cmd == "/system/package/update/check-for-updates" for cmd, _ in fake.calls)
        c.post("/routeros/action", data={"action": "install", "nas": nid, "confirm": "1"})
        c.post("/routeros/action", data={"action": "firmware", "nas": nid, "confirm": "1"})
        assert any(cmd == "/system/package/update/install" for cmd, _ in fake.calls)
        assert any(cmd == "/system/routerboard/upgrade" for cmd, _ in fake.calls)
        c.post("/routeros/delete-npk", data={"npk": "routeros-7.16.2-mipsbe.npk"})
        assert [p["name"] for p in rosadmin.list_npk()] == ["routeros-7.16.2-arm64.npk"]
        # manual purge before a date
        c.post("/data/settings", data={"months": "12"})
        day = (now() - timedelta(days=50)).date()
        from app import jalali
        r = c.post("/data/purge", data={"before": jalali.to_jalali_str(day), "groups": ["usage"]})
        db = SessionLocal()
        assert db.query(SpeedSample).filter(SpeedSample.user_id == uid).count() == 1     # not confirmed
        c.post("/data/purge", data={"before": jalali.to_jalali_str(day), "groups": ["usage"], "confirm": "1"})
        assert db.query(SpeedSample).filter(SpeedSample.user_id == uid).count() == 0
        assert db.query(SpeedSample).filter(SpeedSample.user_id == kid).count() == 2     # "total" quota kept
        assert retention.keep_months(db) == 12
        assert db.query(AuditLog).filter(AuditLog.action == "حذف سابقه").count() == 1
        db.close()
    db = SessionLocal()
    db.delete(db.get(Nas, nid))
    db.commit()
    db.close()


def test_wan_check_ping_verify_access_control_and_windows_pages(env, monkeypatch):
    import json
    from datetime import datetime, timedelta

    from app import access, jobs, mgmt, pingcheck, routeros, wancheck
    from app.db import (AccessExempt, Alert, Device, IfaceDaily, NetObject, SiteRule, now, set_setting)
    from app.web import app
    from tests.fake_ros import FakeRouter
    db = SessionLocal()
    nas = Nas(name="acl-yazd", ip="10.55.0.1", secret="s", api_user="api", lan_ranges="192.168.50.0/24",
              wan_interfaces="ether1,ether3,Pishgaman,wlan1")
    db.add(nas)
    db.commit()
    nid = nas.id
    # --- WAN: Pishgaman (PPPoE) runs over ether1; wlan1 is a port of the LAN bridge
    snap = {"interfaces": [{"name": n, "type": t} for n, t in (("ether1", "ether"), ("ether3", "ether"),
                                                               ("Pishgaman", "pppoe-out"), ("wlan1", "wlan"),
                                                               ("bridge", "bridge"))],
            "addresses": [{"address": "192.168.50.1/24", "interface": "bridge"},
                          {"address": "5.5.5.2/30", "interface": "ether3"}],
            "wan": {"children": [{"name": "Pishgaman", "parent": "ether1"}],
                    "bridge_ports": [{"interface": "wlan1", "bridge": "bridge"}],
                    "routes": [{"gateway": "Pishgaman", "immediate-gw": "Pishgaman", "gateway-status": ""},
                               {"gateway": "5.5.5.1", "immediate-gw": "5.5.5.1%ether3", "gateway-status": ""}]}}
    nas.iface_facts = json.dumps(wancheck.learn(nas, snap))
    day = now().date()
    for name, rx, tx in (("ether1", 10_000, 2_000), ("Pishgaman", 9_000, 1_800), ("ether3", 5_000, 1_000),
                         ("wlan1", 100, 900_000)):
        db.add(IfaceDaily(nas_id=nid, iface=name, day=day, rx=rx, tx=tx))
    db.commit()
    start = datetime.combine(day, datetime.min.time())
    rt = mgmt.router_internet(db, start, start + timedelta(days=1), [nas])
    row = rt["rows"][0]
    assert row["wan_rx"] == 14_000 and row["wan_tx"] == 2_800               # Pishgaman + ether3 only
    counted = {i["name"]: i["counted"] for i in row["ifaces"]}
    assert counted == {"ether1": False, "ether3": True, "Pishgaman": True, "wlan1": False}
    assert set(row["suggested"]) == {"Pishgaman", "ether3"}
    # --- online devices: a static ARP entry alone is a weak sign; a device failing the ping check drops off
    weak = {"arp": [{"mac-address": "AA:00:00:00:00:01", "address": "192.168.50.11", "status": "permanent"}],
            "hotspot_hosts": [{"mac-address": "AA:00:00:00:00:02", "address": "192.168.50.12"}],
            "leases": [], "hotspot_active": [], "ppp_active": []}
    assert not jobs._arp_live(weak["arp"][0])
    assert jobs._arp_live({"mac-address": "AA:00:00:00:00:09", "address": "1.1.1.1", "status": "reachable"})
    ts = now()
    pingcheck.DEAD.clear()
    assert jobs._record_devices(db, nas, weak, ts) == 1                    # the hotspot host (weak, not dead yet)
    db.commit()
    monkeypatch.setattr(pingcheck, "check_router", lambda n, targets: {m: False for m, ip in targets})
    pingcheck.run_once()
    db.expire_all()
    d = db.query(Device).filter(Device.mac == "AA:00:00:00:00:02").one()
    assert d.ping_ok is False and pingcheck.is_dead(d.mac)
    assert pingcheck.online_filter(db.query(Device).filter(Device.mac == d.mac), now()).count() == 0
    assert jobs._record_devices(db, nas, weak, now()) == 0                 # weak sign of a silent device: ignored
    strong = dict(weak, hotspot_hosts=[{"mac-address": "AA:00:00:00:00:02", "address": "192.168.50.12",
                                         "host-dead-time": "5s"}])
    assert jobs._record_devices(db, nas, strong, now()) == 1 and not pingcheck.is_dead(d.mac)
    db.commit()
    # --- access control written to the router
    d.device_type = "phone"
    db.add(NetObject(nas_id=nid, kind="nvr", name="NVR یزد", addresses="192.168.50.200", sensitive=True,
                     sensitive_action="block", allowed_from="192.168.50.15"))
    db.add(NetObject(nas_id=nid, kind="server", name="اتوماسیون", addresses="192.168.50.5", users_allowed=True))
    db.add(SiteRule(pattern="instagram.com", action="block"))
    db.add(AccessExempt(kind="ip", value="192.168.50.40"))
    set_setting(db, "acl_internal_restrict", "1")
    set_setting(db, "public_url", "http://192.168.88.10:9090")
    db.commit()
    fake = FakeRouter()
    fake.store["ip/firewall/filter"] = [{".id": "*A", "chain": "forward", "action": "fasttrack-connection"}]
    monkeypatch.setattr(routeros, "api", fake.conn)
    access.invalidate()
    plan = access.desired(db, nas)
    assert plan["lists"]["macc-block"] == ["instagram.com"] and "192.168.50.5" in plan["lists"]["macc-int-ok"]
    assert "192.168.50.40" in plan["lists"][access.FREE] and "192.168.50.12" in plan["lists"][access.FREE]
    assert {"mac-address": "AA:00:00:00:00:02", "type": "bypassed", "comment": "macc exempt"} in plan["bindings"]
    assert "/ip firewall filter add" in access.script(plan)
    print(access.apply(db, nas))
    rules = fake.store["ip/firewall/filter"]
    assert rules[-1][".id"] == "*A" and all(r["comment"].startswith("macc") for r in rules[:-1])
    n_rules = len(rules) - 1
    assert any(r.get("action") == "drop" and r.get("dst-address-list") == f"macc-s{plan_obj}" for r in rules
               for plan_obj in [db.query(NetObject).filter(NetObject.name == "NVR یزد").one().id])
    assert fake.store["ip/firewall/nat"][0]["to-addresses"] == "192.168.88.10"
    assert {i["address"] for i in fake.store["ip/firewall/address-list"] if i["list"] == "macc-block"} == {"instagram.com"}
    # applying again changes nothing twice; removing leaves the router's own rule
    access.apply(db, nas)
    assert len([r for r in fake.store["ip/firewall/filter"] if r.get("comment", "").startswith("macc")]) == n_rules
    access.apply(db, nas, remove_only=True)
    assert [r[".id"] for r in fake.store["ip/firewall/filter"]] == ["*A"]
    assert not [i for i in fake.store["ip/firewall/address-list"] if i["list"].startswith("macc-")]
    # router caught an access to the NVR
    oid = db.query(NetObject).filter(NetObject.name == "NVR یزد").one().id
    access.violations(db, nas, [{"list": f"macc-v{oid}", "address": "192.168.50.77"}])
    db.commit()
    assert db.query(Alert).filter(Alert.title.like("دسترسی غیرمجاز به NVR یزد از 192.168.50.77%")).count() == 1
    access.violations(db, nas, [{"list": f"macc-v{oid}", "address": "192.168.50.77"}])
    db.commit()
    assert db.query(Alert).filter(Alert.title.like("دسترسی غیرمجاز به NVR یزد%")).count() == 1
    # NetFlow sees it too (allowed source: no alarm)
    access.flow_alarm(SessionLocal, nid, "192.168.50.15", {"id": oid, "name": "NVR یزد", "allowed_from": "192.168.50.15"}, now())
    access.flow_alarm(SessionLocal, nid, "192.168.50.16", {"id": oid, "name": "NVR یزد", "allowed_from": "192.168.50.15"}, now())
    assert db.query(Alert).filter(Alert.title == "دسترسی غیرمجاز به NVR یزد از 192.168.50.16").count() == 1
    db.close()
    with TestClient(app) as c:
        r = c.get("/", headers={"host": "www.instagram.com"})
        assert r.status_code == 403 and "مسدود" in r.text and "www.instagram.com" in r.text
        _login(c, "admin", "Secret12345")
        r = c.get("/access")
        assert r.status_code == 200 and "instagram.com" in r.text and "NVR یزد" in r.text
        assert c.get(f"/access/preview/{nid}").status_code == 200
        c.post("/access/sites/add", data={"patterns": "https://www.YouTube.com/watch\nbad..x", "action": "block"})
        c.post("/access/exempt/add", data={"values": "aa-bb-cc-dd-ee-ff 192.168.50.41"})
        assert c.get("/winlogons?q=nobody").status_code == 200
        r = c.get("/messages")
        assert r.status_code == 200 and "AllowRemoteRPC" in r.text
        c.post("/messages/send", data={"mode": "all", "text": "سلام"})
        r = c.get("/devices")
        assert r.status_code == 200 and "پینگ" in r.text
        r = c.get("/mgmt")
        assert r.status_code == 200
    db = SessionLocal()
    assert db.query(SiteRule).filter(SiteRule.pattern == "youtube.com").count() == 1
    assert db.query(AccessExempt).filter(AccessExempt.value.in_(["AA:BB:CC:DD:EE:FF", "192.168.50.41"])).count() == 2
    for m in (SiteRule, AccessExempt):
        db.query(m).delete()
    db.query(NetObject).filter(NetObject.nas_id == nid).delete()
    db.query(IfaceDaily).filter(IfaceDaily.nas_id == nid).delete()
    set_setting(db, "acl_internal_restrict", "0")
    db.delete(db.get(Nas, nid))
    db.commit()
    db.close()
    pingcheck.DEAD.clear()
    access.invalidate()
