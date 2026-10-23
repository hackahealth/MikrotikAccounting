"""ساخت فایل Excel (راست‌به‌چپ) برای گزارش مدیریتی و ساعات اتصال."""
import io
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import jalali

GB = 1024 ** 3
HEAD_FILL = PatternFill("solid", fgColor="1F3A68")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Tahoma")
BODY_FONT = Font(name="Tahoma")
TITLE_FONT = Font(bold=True, size=14, name="Tahoma")
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def gb(n):
    return round((n or 0) / GB, 3)


def hours(sec):
    return round((sec or 0) / 3600, 2)


def hhmm(dt):
    return dt.strftime("%H:%M") if dt else ""


def sheet(wb, title, headers, rows, widths=None, formats=None, caption=None):
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.rightToLeft = True
    r0 = 1
    if caption:
        ws.cell(row=1, column=1, value=caption).font = TITLE_FONT
        r0 = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY_FONT, BORDER
            if formats and formats.get(c):
                cell.number_format = formats[c]
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 16)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws, r0


def minutes(sec):
    return round((sec or 0) / 60, 1)


def call_sheets(wb, c, caption=None, recent=None, pbx_names=None):
    """Sheets of the call report (pbx.build)."""
    st = c["stats"]
    sheet(wb, "خلاصه تماس‌ها", ["شاخص", "مقدار"], [
        ("تماس ورودی", st["in"]["count"]), ("ورودی پاسخ‌داده", st["in"]["answered"]),
        ("ورودی بی‌پاسخ", st["in"]["count"] - st["in"]["answered"]), ("تماس خروجی", st["out"]["count"]),
        ("خروجی پاسخ‌داده", st["out"]["answered"]), ("تماس داخلی", st["internal"]["count"]),
        ("مجموع مکالمه (دقیقه)", minutes(c["talk"])), ("میانگین مکالمه (ثانیه)", c["avg_talk"]),
        ("میانگین انتظار تا پاسخ ورودی (ثانیه)", c["avg_ring"]),
        ("بی‌پاسخ‌های بدون تماس برگشت", c["missed_count"])], {1: 34, 2: 16}, caption=caption)
    sheet(wb, "تماس افراد", ["کاربر", "نام", "دفتر", "واحد", "داخلی", "ورودی", "ورودی پاسخ‌داده", "بی‌پاسخ",
                             "خروجی", "خروجی پاسخ‌داده", "داخلی", "مکالمه (دقیقه)", "میانگین مکالمه (ثانیه)"],
          [(r["username"], r["full_name"], r["office"], r["department"], r["exts"], r["in"], r["in_answered"],
            r["missed"], r["out"], r["out_answered"], r["internal"], minutes(r["talk"]), r["avg_talk"])
           for r in c["people"]], {2: 22}, {12: "0.0"})
    sheet(wb, "ساعات تماس", ["روز"] + [f"{h:02d}" for h in range(24)],
          [(day,) + tuple(row) for day, row in zip(["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه", "پنجشنبه",
                                                     "جمعه"], c["heat"])], {1: 12})
    sheet(wb, "مقاصد پرتماس", ["شماره", "نوع", "تعداد", "مکالمه (دقیقه)"],
          [(d["number"], d["category"], d["count"], minutes(d["talk"])) for d in c["destinations"]], {1: 18},
          {4: "0.0"})
    sheet(wb, "تماس‌گیرندگان", ["شماره", "تعداد", "بی‌پاسخ", "مکالمه (دقیقه)"],
          [(d["number"], d["count"], d["missed"], minutes(d["talk"])) for d in c["callers"]], {1: 18}, {4: "0.0"})
    sheet(wb, "بی‌پاسخ بدون برگشت", ["شماره", "دفعات", "اولین", "آخرین", "داخلی‌های زنگ‌خورده"],
          [(m["number"], m["count"], jalali.to_jalali_str(m["first"], True), jalali.to_jalali_str(m["last"], True),
            m["exts"]) for m in c["missed"]], {1: 18, 3: 18, 4: 18, 5: 22})
    if c.get("queues"):
        sheet(wb, "صف‌ها", ["صف", "تماس", "پاسخ اپراتور", "رهاشده", "میانگین انتظار (ثانیه)", "بیشترین انتظار (ثانیه)",
                            "انتظار تا رها کردن (ثانیه)", "سطح سرویس ٪", "مکالمه (دقیقه)"],
              [(q["queue"], q["offered"], q["answered"], q["abandoned"], q["avg_wait"], q["max_wait"], q["avg_abandon"],
                q["sl_pct"], minutes(q["talk"])) for q in c["queues"]], {1: 12}, {9: "0.0"})
    if recent is not None:
        names = pbx_names or {}
        sheet(wb, "ریز تماس‌ها", ["زمان", "سرور", "جهت", "داخلی", "طرف مقابل", "وضعیت", "مدت کل (ثانیه)",
                                  "مکالمه (ثانیه)"],
              [(jalali.to_jalali_str(r.ts, True), names.get(r.pbx_id, ""), {"in": "ورودی", "out": "خروجی",
                "internal": "داخلی"}.get(r.direction, "نامشخص"), r.extension, r.other, r.disposition, r.duration,
                r.billsec) for r in recent], {1: 18, 5: 18})


def calls_xlsx(c, caption, recent, pbx_names):
    wb = Workbook()
    wb.remove(wb.active)
    call_sheets(wb, c, caption, recent, pbx_names)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def management_xlsx(rep, title_extra="", calls=None):
    k = rep["kpi"]
    wb = Workbook()
    wb.remove(wb.active)
    period = f"{jalali.to_jalali_str(rep['start'])} تا {jalali.to_jalali_str(rep['end'] - timedelta(days=1))}"

    def chg(v):
        return "-" if v is None else f"{v:+.1f}٪"

    summary = [
        ("دوره گزارش", period),
        ("فیلتر", title_extra or "همه کاربران"),
        ("ساعات کاری", rep["work_hours"]),
        ("مجموع مصرف (GB)", gb(k["total"])),
        ("تغییر نسبت به دوره قبل", chg(k["total_change"])),
        ("مصرف اینترنت کاربران (GB)", gb(k.get("internet", 0))),
        ("مصرف داخل شبکه (GB)", gb(k.get("internal", 0))),
        ("مصرف تانل - شبکه دفاتر دیگر (GB)", gb(k.get("tunnel", 0))),
        ("سهم اینترنت از کل مصرف", f"{k.get('internet_pct', 0)}٪"),
        ("کاربران فعال", k["users"]),
        ("تغییر کاربران فعال", chg(k["users_change"])),
        ("میانگین مصرف هر کاربر (GB)", gb(k["per_user"])),
        ("مجموع ساعات آنلاین", k["online_h"]),
        ("میانگین ساعات آنلاین هر کاربر", k["online_per_user_h"]),
        ("تعداد اتصال", k["sessions"]),
        ("مصرف در ساعات کاری (GB)", gb(k["work_bytes"])),
        ("مصرف خارج از ساعات کاری (GB)", gb(k["off_bytes"])),
        ("سهم ساعات کاری از مصرف", f"{k['work_pct']}٪"),
        ("ترافیک تفریحی در ساعات کاری (GB)", gb(k["leisure_work"])),
        ("سهم ترافیک تفریحی از ترافیک ساعات کاری", "-" if k["leisure_work_pct"] is None else f"{k['leisure_work_pct']}٪"),
        ("ساعت اوج مصرف", "-" if k["peak_hour"] is None else f"{k['peak_hour']:02d}:00"),
        ("روز اوج مصرف", "-" if not k["peak_day"] else jalali.to_jalali_str(k["peak_day"]["day"])),
        ("ورود ناموفق", k["rejects"]),
        ("کاربران با حجم تمام‌شده", k["quota_users"]),
    ]
    rt = rep.get("routers")
    if rt and rt["rows"]:
        summary[3:3] = [("کل اینترنت واقعی - اینترفیس WAN روترها (GB)", gb(rt["total"]["wan"])),
                        ("تانل - اینترفیس‌های تانل (GB)", gb(rt["total"]["tunnel"])),
                        ("اینترنت منهای تانل (GB)", gb(rt["total"]["pure"]))]
    sheet(wb, "خلاصه", ["شاخص", "مقدار"], summary, {1: 42, 2: 40}, caption="گزارش مدیریتی مصرف اینترنت")
    if rt and rt["rows"]:
        sheet(wb, "روترها", ["روتر", "اینترفیس WAN", "اینترفیس تانل", "کل اینترنت (GB)", "دانلود (GB)", "آپلود (GB)",
                             "تانل (GB)", "اینترنت منهای تانل (GB)"],
              [(r["name"], ", ".join(r["wan_ifs"]), ", ".join(r["tun_ifs"]), gb(r["wan"]), gb(r["wan_rx"]),
                gb(r["wan_tx"]), gb(r["tunnel"]), gb(r["pure"])) for r in rt["rows"]],
              {1: 20, 2: 20, 3: 24}, caption="از شمارنده بایت اینترفیس‌های روتر، مستقل از کاربران")

    ws, r0 = sheet(wb, "روند روزانه", ["تاریخ", "دانلود (GB)", "آپلود (GB)", "مجموع (GB)", "اینترنت (GB)",
                                       "داخل شبکه (GB)", "تانل (GB)"],
                   [(jalali.to_jalali_str(d["day"]), gb(d["down"]), gb(d["up"]), gb(d["total"]),
                     gb(d.get("internet", d["total"])), gb(d.get("internal", 0)), gb(d.get("tunnel", 0)))
                    for d in rep["daily"]],
                   formats={2: "0.000", 3: "0.000", 4: "0.000", 5: "0.000", 6: "0.000", 7: "0.000"})
    if rep["daily"]:
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
        ch.title, ch.y_axis.title, ch.height, ch.width = "مصرف روزانه: اینترنت / داخل شبکه / تانل (GB)", "GB", 8, 22
        ch.add_data(Reference(ws, min_col=5, max_col=7, min_row=r0, max_row=r0 + len(rep["daily"])),
                    titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + len(rep["daily"])))
        ws.add_chart(ch, "H2")
    sheet(wb, "مصرف روزانه کاربران", ["تاریخ", "کاربر", "نام", "دفتر", "واحد", "اینترنت (GB)", "داخل شبکه (GB)",
                                     "مجموع (GB)", "دانلود اینترنت (GB)", "آپلود اینترنت (GB)",
                                     "دانلود داخلی (GB)", "آپلود داخلی (GB)", "تانل (GB)", "دانلود تانل (GB)",
                                     "آپلود تانل (GB)"],
          [(jalali.to_jalali_str(r["day"]), r.get("username", ""), r.get("full_name", ""), r.get("office", ""),
            r.get("department", ""), gb(r["internet"]), gb(r["internal"]), gb(r["total"]), gb(r["internet_down"]),
            gb(r["internet_up"]), gb(r["internal_down"]), gb(r["internal_up"]), gb(r.get("tunnel", 0)),
            gb(r.get("tunnel_down", 0)), gb(r.get("tunnel_up", 0))) for r in rep.get("user_daily", [])],
          {3: 22}, {c: "0.000" for c in range(6, 16)})

    sheet(wb, "گروه‌ها", ["گروه", "کاربران", "مصرف (GB)", "سهم ٪", "میانگین هر کاربر (GB)", "ساعات آنلاین"],
          [(g["name"], g["users"], gb(g["bytes"]), g["share"], gb(g["per_user"]), g["online_h"]) for g in rep["by_group"]],
          {1: 24}, {3: "0.000", 5: "0.000"})
    for title, key in (("دفترها", "by_office"), ("واحدها", "by_department")):
        if rep.get(key):
            sheet(wb, title, ["دفتر" if key == "by_office" else "واحد", "کاربران", "مصرف (GB)", "سهم ٪",
                              "میانگین هر کاربر (GB)", "ساعات آنلاین"],
                  [(g["name"], g["users"], gb(g["bytes"]), g["share"], gb(g["per_user"]), g["online_h"])
                   for g in rep[key]], {1: 24}, {3: "0.000", 5: "0.000"})
    sheet(wb, "کاربران", ["رتبه", "کاربر", "نام", "دفتر", "واحد", "گروه", "دانلود (GB)", "آپلود (GB)", "مجموع (GB)", "سهم ٪",
                         "ساعات آنلاین", "بیشترین نوع استفاده", "تفریحی در ساعت کاری ٪", "اینترنت (GB)",
                         "داخل شبکه (GB)", "تانل (GB)"],
          [(i, u["username"], u["full_name"], u.get("office", ""), u.get("department", ""), u["group"], gb(u["down"]),
            gb(u["up"]), gb(u["total"]), u["share"], u["online_h"], u["top_category"],
            "" if u["leisure_work_pct"] is None else u["leisure_work_pct"], gb(u.get("internet", u["total"])),
            gb(u.get("internal", 0)), gb(u.get("tunnel", 0)))
           for i, u in enumerate(rep["top_users"], 1)],
          {1: 7, 3: 22, 12: 20}, {7: "0.000", 8: "0.000", 9: "0.000", 14: "0.000", 15: "0.000", 16: "0.000"})
    ws, r0 = sheet(wb, "نوع استفاده", ["نوع استفاده", "مصرف (GB)", "سهم ٪", "در ساعات کاری (GB)"],
                   [(c["name"], gb(c["bytes"]), c["pct"], gb(c["work"])) for c in rep["categories"]],
                   {1: 26}, {2: "0.000", 4: "0.000"})
    if rep["categories"]:
        pie = PieChart()
        pie.title, pie.height, pie.width = "سهم نوع استفاده", 9, 14
        pie.add_data(Reference(ws, min_col=2, min_row=r0, max_row=r0 + len(rep["categories"])), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + len(rep["categories"])))
        ws.add_chart(pie, "F2")
    sheet(wb, "ساعات شبانه‌روز", ["ساعت", "مصرف (GB)"],
          [(f"{h:02d}:00", gb(b)) for h, b in enumerate(rep["by_hour"])], formats={2: "0.000"})
    sheet(wb, "سایت‌ها", ["سایت / سرویس", "نوع", "مصرف (GB)", "تعداد کاربر"],
          [(d["domain"], d["category"], gb(d["bytes"]), d["users"]) for d in rep["domains"]], {1: 34}, {3: "0.000"})
    sheet(wb, "دستگاه و سرویس", ["نوع", "تعداد اتصال", "تعداد کاربر", "مصرف (GB)"],
          [(d["name"], d["sessions"], d["users"], gb(d["bytes"])) for d in rep["devices"]] +
          [(s["name"], s["sessions"], "", gb(s["bytes"])) for s in rep["services"]], formats={4: "0.000"})
    sheet(wb, "ورودهای ناموفق", ["علت", "تعداد", "تعداد کاربر"],
          [(r["reason"], r["count"], r["users"]) for r in rep["rejects"]], {1: 40})
    if calls and calls["calls"]:
        call_sheets(wb, calls)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def attendance_xlsx(rows, caption):
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "ساعات اتصال", ["تاریخ", "روز کاری", "کاربر", "نام", "کد پرسنلی", "دفتر", "واحد", "گروه",
                             "اولین اتصال", "آخرین اتصال", "ساعات آنلاین", "در ساعات کاری", "خارج از ساعات کاری",
                             "تعداد اتصال", "مصرف (GB)", "سرویس", "دستگاه", "اینترنت (GB)", "داخل شبکه (GB)",
                             "تانل (GB)"],
          [(jalali.to_jalali_str(r["day"]), "بله" if r["work_day"] else "خیر", r["username"], r["full_name"],
            r.get("personnel_code", ""), r.get("office", ""), r.get("department", ""), r["group"],
            hhmm(r["first"]), "آنلاین" if r["still_online"] else hhmm(r["last"]), hours(r["online"]),
            hours(r["inside"]), hours(r["outside"]), r["sessions"], gb(r["bytes"]), r["services"], r["devices"],
            gb(r.get("internet", r["bytes"])), gb(r.get("internal", 0)), gb(r.get("tunnel", 0)))
           for r in rows], {4: 22, 16: 16, 17: 18},
          {11: "0.00", 12: "0.00", 13: "0.00", 15: "0.000", 18: "0.000", 19: "0.000", 20: "0.000"},
          caption=caption)
    # per-user summary
    agg = {}
    for r in rows:
        a = agg.setdefault(r["user_id"], {"u": r, "days": 0, "online": 0, "inside": 0, "outside": 0, "bytes": 0,
                                          "internet": 0, "internal": 0, "tunnel": 0, "firsts": []})
        a["days"] += 1
        for k in ("online", "inside", "outside", "bytes"):
            a[k] += r[k]
        a["internet"] += r.get("internet", r["bytes"])
        a["internal"] += r.get("internal", 0)
        a["tunnel"] += r.get("tunnel", 0)
        if r["work_day"]:
            a["firsts"].append(r["first"].hour * 60 + r["first"].minute)
    summary = []
    for a in sorted(agg.values(), key=lambda x: x["u"]["username"]):
        avg_first = sum(a["firsts"]) // len(a["firsts"]) if a["firsts"] else None
        summary.append((a["u"]["username"], a["u"]["full_name"], a["u"].get("office", ""),
                        a["u"].get("department", ""), a["u"]["group"], a["days"], hours(a["online"]),
                        hours(a["inside"]), hours(a["outside"]),
                        "" if avg_first is None else f"{avg_first // 60:02d}:{avg_first % 60:02d}", gb(a["bytes"]),
                        gb(a["internet"]), gb(a["internal"]), gb(a["tunnel"])))
    sheet(wb, "خلاصه کارکنان", ["کاربر", "نام", "دفتر", "واحد", "گروه", "روزهای دارای اتصال", "ساعات آنلاین",
                               "در ساعات کاری", "خارج از ساعات کاری", "میانگین اولین اتصال (روز کاری)", "مصرف (GB)",
                               "اینترنت (GB)", "داخل شبکه (GB)", "تانل (GB)"],
          summary, {2: 22, 10: 22}, {7: "0.00", 8: "0.00", 9: "0.00", 11: "0.000", 12: "0.000", 13: "0.000",
                                     14: "0.000"})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def simple_xlsx(title, headers, rows, caption=None, widths=None):
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, title, headers, rows, widths or {2: 22, 4: 19, 7: 18, 8: 22, 9: 18, 11: 26}, caption=caption)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
